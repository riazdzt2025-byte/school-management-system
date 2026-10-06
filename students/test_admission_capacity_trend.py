"""OF-04 (prompt 12) — capacity vs enrolled, and the admission trend.

The funnel report itself is pinned in ``test_admission_funnel_report.py``; this
file covers what owner decisions in ``prompt-12-of-04-admission-reports.md`` §8
added beside it:

* capacity vs enrolled per class/section (``SectionCapacity`` against the same
  ACTIVE students the admission seat check counts),
* the date-bucketed trend of submissions, payment approvals and enrolments —
  table + CSS bars, no new chart dependency, and
* the Accounts payment-stage slice: the page, the export and the trend all show
  the payment stages and nothing office-side.

Only synthetic data in Django's disposable test database.
"""
import re
import unittest
from datetime import date, datetime, timedelta
from datetime import timezone as dt_timezone
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

try:
    from openpyxl import load_workbook
except ModuleNotFoundError:  # pragma: no cover - openpyxl is a hard dependency
    load_workbook = None

from .models import AdmissionApplication, Institution, InstitutionAccess, SectionCapacity, Student
from .test_admission_funnel_report import _count_on_page, _funnel_sheet
from .views import CAPACITY_DISPLAY_LIMIT, TREND_DISPLAY_LIMIT

requires_openpyxl = unittest.skipIf(load_workbook is None, 'openpyxl is required for the Excel export tests')

PAGE = 'admission_funnel_report'
EXPORT = 'admission_funnel_export'
LABELS = dict(AdmissionApplication.STATUS_CHOICES)


def _utc(year, month, day, hour=9, minute=0):
    """An aware datetime in UTC — the fixture writes its own timestamps."""
    return datetime(year, month, day, hour, minute, tzinfo=dt_timezone.utc)


def _capacity_row(response, institution, admission_class, section):
    for row in response.context['capacity_rows']:
        if (row['institution'], row['admission_class'], row['section']) == (
            institution, admission_class, section,
        ):
            return row
    raise AssertionError(
        f'no capacity row for {institution!r} class {admission_class!r} section {section!r}'
    )


def _trend_values(response):
    """{bucket label: [cell values]} in the order the table renders them."""
    return {
        row['label']: [cell['value'] for cell in row['cells']]
        for row in response.context['trend_rows']
    }


def _trend_sheet_rows(workbook):
    return list(workbook['Trend'].iter_rows(values_only=True))


class AdmissionReportFixture(TestCase):
    """Shared fixture: two institutions, the three roles, and a small cohort."""

    def setUp(self):
        self.institution_a = Institution.objects.create(name='OF04 A', classes='6,7')
        self.institution_b = Institution.objects.create(name='OF04 B', classes='6,7')

        self.admin = get_user_model().objects.create_superuser(
            username='of04-admin', password='password', email='of04@example.com',
        )
        self.office_clerk = self._clerk('of04-office', 'Office')
        self.accounts_clerk = self._clerk('of04-accounts', 'Accounts')

    def _clerk(self, username, department):
        user = get_user_model().objects.create_user(username=username, password='password')
        InstitutionAccess.objects.create(
            user=user, institution=self.institution_a, department=department,
        )
        user.user_permissions.add(
            Permission.objects.get(
                content_type=ContentType.objects.get_for_model(AdmissionApplication),
                codename='view_admissionapplication',
            )
        )
        return user

    def _student(self, institution, admission_class, section, status='ACTIVE'):
        return Student.objects.create(
            institution=institution, name=f'OF04 {admission_class}{section} {status}',
            admission_class=admission_class, section=section, status=status,
            guardian_contact_no='01900000000',
        )

    def _application(self, institution, status='SUBMITTED', submitted_at=None,
                     account_action_at=None, name=None):
        """An application whose timestamps the fixture controls.

        ``submitted_at`` is ``auto_now_add`` and ``account_action_at`` is only
        ever set by the accounts view, so a fixture that needs a specific moment
        writes it through the queryset (which skips ``auto_now_add``).
        """
        application = AdmissionApplication.objects.create(
            institution=institution,
            applicant_name=name or f'{status} applicant',
            guardian_name='Guardian', guardian_contact_no='01900000000',
            requested_class='6', session='2026-2027', status=status,
        )
        update = {}
        if submitted_at is not None:
            update['submitted_at'] = submitted_at
        if account_action_at is not None:
            update['account_action_at'] = account_action_at
        if update:
            AdmissionApplication.objects.filter(pk=application.pk).update(**update)
        return application

    def _login(self, user, institution=None, department='Office'):
        self.client.force_login(user)
        session = self.client.session
        if institution is not None:
            session['selected_institution_id'] = str(institution.pk)
        session['selected_department'] = department
        session.save()

    def _page_queries(self):
        """How many queries one page render costs (aggregate-query evidence)."""
        with CaptureQueriesContext(connection) as captured:
            response = self.client.get(reverse(PAGE))
            self.assertEqual(response.status_code, 200)
        return len(captured)


class CapacityVsEnrolledTests(AdmissionReportFixture):
    """Seat limit vs the students actually sitting in each class/section."""

    def setUp(self):
        super().setUp()
        # A / 6 / A: three active students (plus one discontinued, never counted)
        # against a limit of 5. A / 7 / B holds students with no limit row at
        # all, A / 6 / B has a zero limit, and B / 6 / A is over its limit.
        for _ in range(3):
            self._student(self.institution_a, '6', 'A')
        self._student(self.institution_a, '6', 'A', status='DISCONTINUED')
        for _ in range(2):
            self._student(self.institution_a, '7', 'B')
        for _ in range(5):
            self._student(self.institution_b, '6', 'A')

        SectionCapacity.objects.create(
            institution=self.institution_a, admission_class='6', section='A', capacity=5,
        )
        SectionCapacity.objects.create(
            institution=self.institution_a, admission_class='6', section='B', capacity=0,
        )
        SectionCapacity.objects.create(
            institution=self.institution_b, admission_class='6', section='A', capacity=2,
        )

    def test_capacity_row_shows_limit_seated_free_and_utilization(self):
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))

        row = _capacity_row(response, 'OF04 A', '6', 'A')
        self.assertTrue(row['has_limit'])
        self.assertEqual(row['capacity'], 5)
        self.assertEqual(row['seated'], 3)  # the discontinued student is out
        self.assertEqual(row['free'], 2)
        self.assertEqual(row['utilization'], 60.0)
        self.assertEqual(row['utilization_display'], '60.0%')
        self.assertFalse(row['over'])
        self.assertEqual(row['bar_percent'], 60.0)

        self.assertContains(response, 'Seat capacity vs enrolled students (ACTIVE)')
        self.assertContains(response, '60.0%')

    def test_section_with_students_but_no_limit_row_is_listed_as_no_limit(self):
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))

        row = _capacity_row(response, 'OF04 A', '7', 'B')
        self.assertFalse(row['has_limit'])
        self.assertEqual(row['capacity_display'], 'No limit')
        self.assertEqual(row['seated'], 2)
        self.assertEqual(row['free_display'], '—')
        self.assertEqual(row['utilization_display'], '—')
        self.assertEqual(row['bar_percent'], 0.0)
        self.assertContains(response, 'No limit configured')

    def test_zero_capacity_and_over_capacity_are_not_division_by_zero(self):
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))
        self.assertEqual(response.status_code, 200)

        zero = _capacity_row(response, 'OF04 A', '6', 'B')
        self.assertEqual(zero['capacity'], 0)
        self.assertEqual(zero['seated'], 0)
        self.assertIsNone(zero['utilization'])  # 0/0 stays '—'
        self.assertEqual(zero['free'], 0)
        self.assertFalse(zero['over'])

        over = _capacity_row(response, 'OF04 B', '6', 'A')
        self.assertTrue(over['over'])
        self.assertEqual(over['over_by'], 3)
        self.assertEqual(over['free'], -3)
        self.assertEqual(over['utilization'], 250.0)
        self.assertEqual(over['bar_percent'], 100.0)  # the bar never leaves its track
        self.assertContains(response, 'Over by 3')

    def test_capacity_totals_only_promise_free_seats_where_a_limit_exists(self):
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))
        totals = response.context['capacity_totals']

        self.assertEqual(totals['sections'], 4)
        self.assertEqual(totals['limited_sections'], 3)
        self.assertEqual(totals['no_limit_sections'], 1)
        self.assertEqual(totals['capacity'], 7)  # 5 + 0 + 2
        self.assertEqual(totals['seated'], 10)   # 3 + 0 + 2 + 5
        self.assertEqual(totals['free'], 7 - (3 + 0 + 5))  # limited sections only
        self.assertEqual(totals['over_sections'], 1)

    def test_capacity_is_institution_scoped_and_the_query_param_cannot_widen_it(self):
        self._login(self.office_clerk, institution=self.institution_a)
        response = self.client.get(reverse(PAGE))
        body = response.content.decode()
        self.assertIn('OF04 A', body)
        self.assertNotIn('OF04 B', body)

        widened = self.client.get(reverse(PAGE), {'institution': self.institution_b.pk})
        self.assertEqual(widened.status_code, 200)
        self.assertEqual(widened.context['institution'], self.institution_a)
        # A/6/A, A/6/B and A/7/B — B's students and limit stay out.
        self.assertEqual(widened.context['capacity_totals']['sections'], 3)
        for row in widened.context['capacity_rows']:
            self.assertEqual(row['institution'], 'OF04 A')
        self.assertNotContains(widened, 'OF04 B')

    @requires_openpyxl
    def test_capacity_sheet_pins_the_header_and_the_limit_semantics(self):
        self._login(self.admin)
        response = self.client.get(reverse(EXPORT))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        self.assertIn('Capacity vs Enrolled', workbook.sheetnames)

        rows = list(workbook['Capacity vs Enrolled'].iter_rows(values_only=True))
        header = rows.index((
            'Institution', 'Class', 'Section', 'Capacity', 'Enrolled (ACTIVE)',
            'Free seats', 'Utilization %', 'Status',
        ))
        by_key = {(row[0], row[1], row[2]): row for row in rows[header + 1:] if row[0]}
        self.assertEqual(by_key[('OF04 A', '6', 'A')][3:8], (5, 3, 2, 60.0, 'Within limit'))
        self.assertEqual(by_key[('OF04 A', '7', 'B')][3], 'No limit')
        self.assertEqual(by_key[('OF04 A', '7', 'B')][7], 'No limit')
        self.assertEqual(by_key[('OF04 B', '6', 'A')][7], 'Over capacity')
        self.assertEqual(by_key[('OF04 B', '6', 'A')][5], -3)

    @requires_openpyxl
    def test_capacity_page_is_bounded_while_the_sheet_keeps_every_row(self):
        SectionCapacity.objects.bulk_create([
            SectionCapacity(
                institution=self.institution_b, admission_class=f'X{i:03d}', section='', capacity=i + 1,
            )
            for i in range(CAPACITY_DISPLAY_LIMIT + 5)
        ])
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))

        total = response.context['capacity_totals']['sections']
        self.assertEqual(len(response.context['capacity_rows']), CAPACITY_DISPLAY_LIMIT)
        self.assertEqual(response.context['capacity_truncated'], total - CAPACITY_DISPLAY_LIMIT)
        self.assertGreaterEqual(response.context['capacity_truncated'], 5)
        self.assertContains(
            response, f'Showing the first {CAPACITY_DISPLAY_LIMIT} of {total} rows',
        )

        workbook = load_workbook(BytesIO(self.client.get(reverse(EXPORT)).content), data_only=True)
        sheet = workbook['Capacity vs Enrolled']
        sheet_rows = list(sheet.iter_rows(values_only=True))
        header = next(
            index for index, row in enumerate(sheet_rows) if row[:3] == ('Institution', 'Class', 'Section')
        )
        data_rows = [row for row in sheet_rows[header + 1:] if isinstance(row[1], str) and row[1]]
        self.assertEqual(len(data_rows), total)  # the export is never truncated

    def test_capacity_section_adds_no_query_per_row(self):
        self._login(self.admin)
        before = self._page_queries()
        SectionCapacity.objects.bulk_create([
            SectionCapacity(
                institution=self.institution_b, admission_class=f'Q{i:03d}', section='', capacity=10,
            )
            for i in range(60)
        ])
        after = self._page_queries()
        self.assertLessEqual(after, before + 2)

    def test_capacity_section_is_visible_to_accounts_too(self):
        """Seats are not an office-stage secret: the payment step checks them."""
        self._login(self.accounts_clerk, institution=self.institution_a, department='Accounts')
        response = self.client.get(reverse(PAGE))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Seat capacity vs enrolled students (ACTIVE)')
        self.assertEqual(_capacity_row(response, 'OF04 A', '6', 'A')['capacity'], 5)


class AdmissionTrendTests(AdmissionReportFixture):
    """The date-bucketed trend: buckets, boundaries, series and windows."""

    def test_trend_buckets_per_day_in_chronological_order(self):
        self._application(self.institution_a, submitted_at=_utc(2026, 3, 1))
        self._application(self.institution_a, submitted_at=_utc(2026, 3, 1, 12))
        self._application(self.institution_a, submitted_at=_utc(2026, 3, 3))
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))

        self.assertEqual(response.context['trend_granularity'], 'day')
        self.assertEqual(
            [row['label'] for row in response.context['trend_rows']],
            ['2026-03-01', '2026-03-03'],
        )
        first = response.context['trend_rows'][0]
        self.assertEqual([cell['value'] for cell in first['cells']], [2, 0, 0])
        self.assertEqual(first['cells'][0]['percent'], 100.0)
        self.assertContains(response, 'Trend per day — day boundaries in UTC')

    def test_trend_bucket_boundary_follows_the_timezone(self):
        """23:01 and 00:01 in Dhaka — one minute apart, on different days."""
        self._application(self.institution_a, submitted_at=_utc(2026, 1, 1, 17, 1))
        self._application(self.institution_a, submitted_at=_utc(2026, 1, 1, 18, 1))
        self._login(self.admin)

        utc_view = self.client.get(reverse(PAGE))
        self.assertEqual(_trend_values(utc_view), {'2026-01-01': [2, 0, 0]})

        with timezone.override('Asia/Dhaka'):
            dhaka_view = self.client.get(reverse(PAGE))
        self.assertEqual(_trend_values(dhaka_view), {'2026-01-01': [1, 0, 0], '2026-01-02': [1, 0, 0]})
        self.assertContains(dhaka_view, 'day boundaries in Asia/Dhaka')

    def test_trend_series_are_dated_by_their_own_event(self):
        self._application(
            self.institution_a, 'ENROLLED',
            submitted_at=_utc(2026, 4, 1), account_action_at=_utc(2026, 4, 10, 11),
        )
        self._application(
            self.institution_a, 'PAYMENT_APPROVED',
            submitted_at=_utc(2026, 4, 2), account_action_at=_utc(2026, 4, 12, 11),
        )
        self._login(self.admin)
        values = _trend_values(self.client.get(reverse(PAGE)))

        self.assertEqual(values['2026-04-01'], [1, 0, 0])
        self.assertEqual(values['2026-04-02'], [1, 0, 0])
        # Enrolment happened in the payment transaction, so that day carries both.
        self.assertEqual(values['2026-04-10'], [0, 1, 1])
        # Approved but not enrolled yet: the difference the office chases.
        self.assertEqual(values['2026-04-12'], [0, 1, 0])

    def test_trend_undated_row_is_explicit_and_never_dropped(self):
        self._application(
            self.institution_a, 'ENROLLED', submitted_at=_utc(2026, 5, 1), account_action_at=None,
        )
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))
        rows = response.context['trend_rows']

        self.assertEqual(rows[-1]['label'], 'Undated')
        self.assertTrue(rows[-1]['undated'])
        self.assertEqual(rows[-1]['cells'][2]['value'], 1)
        self.assertIsNone(rows[-1]['cells'][0]['display'])
        # The dated buckets are untouched by it.
        self.assertEqual([row['label'] for row in rows[:-1]], ['2026-05-01'])
        self.assertContains(response, 'no accounts action time')

    def test_trend_week_month_and_invalid_bucket(self):
        self._application(self.institution_a, submitted_at=_utc(2026, 1, 1))  # ISO week 1
        self._application(self.institution_a, submitted_at=_utc(2026, 1, 5))  # ISO week 2
        self._login(self.admin)

        week = self.client.get(reverse(PAGE), {'bucket': 'week'})
        self.assertEqual(week.context['trend_granularity'], 'week')
        self.assertEqual([row['label'] for row in week.context['trend_rows']], ['2026-W01', '2026-W02'])
        self.assertEqual(
            [row['start'] for row in week.context['trend_rows']], ['2025-12-29', '2026-01-05'],
        )

        month = self.client.get(reverse(PAGE), {'bucket': 'month'})
        self.assertEqual([row['label'] for row in month.context['trend_rows']], ['2026-01'])
        self.assertEqual(month.context['trend_rows'][0]['cells'][0]['value'], 2)

        junk = self.client.get(reverse(PAGE), {'bucket': 'hour'})
        self.assertEqual(junk.status_code, 200)
        self.assertEqual(junk.context['trend_granularity'], 'day')

    def test_trend_window_is_applied_to_each_series_own_event_date(self):
        """A February payment on a January submission shows up in February."""
        self._application(
            self.institution_a, 'ENROLLED',
            submitted_at=_utc(2026, 1, 5), account_action_at=_utc(2026, 2, 10, 10),
        )
        self._login(self.admin)
        response = self.client.get(reverse(PAGE), {'from': '2026-02-01', 'to': '2026-02-28'})

        self.assertEqual(_trend_values(response), {'2026-02-10': [0, 1, 1]})
        # The funnel half still counts *submissions* in the window — nothing there.
        self.assertEqual(response.context['total_applications'], 0)

    def test_trend_is_institution_scoped_and_cannot_be_widened(self):
        self._application(self.institution_b, submitted_at=_utc(2026, 6, 1))
        self._login(self.office_clerk, institution=self.institution_a)

        self.assertEqual(self.client.get(reverse(PAGE)).context['trend_rows'], [])
        widened = self.client.get(reverse(PAGE), {'institution': self.institution_b.pk})
        self.assertEqual(widened.context['trend_rows'], [])
        self.assertNotContains(widened, 'OF04 B')

    def test_trend_bars_are_css_scaled_to_the_busiest_number(self):
        self._application(self.institution_a, submitted_at=_utc(2026, 10, 1))
        self._application(self.institution_a, submitted_at=_utc(2026, 10, 2))
        self._application(self.institution_a, submitted_at=_utc(2026, 10, 2, 15))
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))

        values = {row['label']: [cell['percent'] for cell in row['cells']]
                  for row in response.context['trend_rows']}
        self.assertEqual(values['2026-10-01'][0], 50.0)   # 1 of the busiest 2
        self.assertEqual(values['2026-10-02'][0], 100.0)

        body = response.content.decode()
        self.assertIn('funnel-bar bar-submitted', body)
        self.assertNotIn('<canvas', body)
        self.assertNotIn('chart.js', body.lower())

    def test_trend_page_is_bounded_and_the_sheet_keeps_every_bucket(self):
        days = TREND_DISPLAY_LIMIT + 5
        for index in range(days):
            self._application(
                self.institution_a, submitted_at=_utc(2026, 1, 1) + timedelta(days=index),
            )
        self._login(self.admin)
        response = self.client.get(reverse(PAGE))

        self.assertEqual(len(response.context['trend_rows']), TREND_DISPLAY_LIMIT)
        self.assertEqual(response.context['trend_truncated'], 5)
        self.assertContains(
            response, f'Showing the most recent {TREND_DISPLAY_LIMIT} buckets',
        )
        newest = date(2026, 1, 1) + timedelta(days=days - 1)
        self.assertEqual(response.context['trend_rows'][-1]['label'], newest.isoformat())

        workbook = load_workbook(BytesIO(self.client.get(reverse(EXPORT)).content), data_only=True)
        labels = [
            row[0] for row in _trend_sheet_rows(workbook)
            if isinstance(row[0], str) and re.match(r'^\d{4}-\d{2}-\d{2}$', row[0])
        ]
        self.assertEqual(len(labels), days)
        self.assertEqual(labels[0], '2026-01-01')
        self.assertEqual(labels[-1], newest.isoformat())

    @requires_openpyxl
    def test_trend_sheet_pins_its_header_meta_and_bucket_size(self):
        self._application(
            self.institution_a, 'ENROLLED',
            submitted_at=_utc(2026, 7, 1), account_action_at=_utc(2026, 7, 2, 10),
        )
        self._login(self.admin)
        workbook = load_workbook(BytesIO(self.client.get(reverse(EXPORT)).content), data_only=True)
        rows = _trend_sheet_rows(workbook)

        self.assertEqual(rows[0][0], 'Admission Trend Report')
        self.assertEqual(rows[1][:2], ('Institution', 'All Institutions'))
        self.assertEqual(rows[2][:2], ('Bucket', 'day'))
        header = next(row for row in rows if row[0] == 'Bucket' and row[1] == 'Bucket start')
        self.assertEqual(
            header,
            ('Bucket', 'Bucket start', 'Submitted', 'Payment approved', 'Enrolled'),
        )
        data = [row for row in rows[rows.index(header) + 1:] if isinstance(row[1], str)]
        self.assertEqual(data[0][:5], ('2026-07-01', '2026-07-01', 1, 0, 0))
        self.assertEqual(data[1][:5], ('2026-07-02', '2026-07-02', 0, 1, 1))

        # ?bucket= travels to the workbook, so the download matches the page.
        monthly = load_workbook(
            BytesIO(self.client.get(reverse(EXPORT), {'bucket': 'month'}).content), data_only=True,
        )
        month_rows = _trend_sheet_rows(monthly)
        self.assertEqual(month_rows[2][:2], ('Bucket', 'month'))
        self.assertIn(('2026-07', '2026-07-01', 1, 1, 1), month_rows)

    def test_trend_table_adds_no_query_per_bucket(self):
        self._login(self.admin)
        for index in range(3):
            self._application(
                self.institution_a, submitted_at=_utc(2026, 9, 1) + timedelta(days=index),
            )
        before = self._page_queries()
        for index in range(3, 60):
            self._application(
                self.institution_a, submitted_at=_utc(2026, 9, 1) + timedelta(days=index),
            )
        after = self._page_queries()
        self.assertLessEqual(after, before + 2)


class AccountsPaymentSliceTests(AdmissionReportFixture):
    """Owner decision (prompt 12 §8): Accounts reads the payment stages only."""

    def setUp(self):
        super().setUp()
        for index, status in enumerate(
            ('SUBMITTED', 'OFFICE_APPROVED', 'ACCOUNT_PENDING', 'PAYMENT_APPROVED', 'ENROLLED', 'REJECTED')
        ):
            self._application(
                self.institution_a, status,
                submitted_at=_utc(2026, 2, 1 + index),
                account_action_at=(
                    _utc(2026, 3, 1 + index)
                    if status in ('PAYMENT_APPROVED', 'ENROLLED') else None
                ),
            )

    def test_accounts_page_shows_only_the_payment_stages(self):
        self._login(self.accounts_clerk, institution=self.institution_a, department='Accounts')
        response = self.client.get(reverse(PAGE))

        self.assertEqual(
            response.context['stage_keys'], ['ACCOUNT_PENDING', 'PAYMENT_APPROVED', 'ENROLLED'],
        )
        # One application sits in each of the three slices.
        self.assertEqual(response.context['total_applications'], 3)
        body = response.content.decode()
        self.assertIn('Accounts view — payment stages only', body)
        self.assertNotIn(f'<td>{LABELS["SUBMITTED"]}</td>', body)
        self.assertNotIn(f'<td>{LABELS["OFFICE_APPROVED"]}</td>', body)
        self.assertNotIn(f'Rejected</div>', body)
        self.assertEqual(_count_on_page(response, 'ACCOUNT_PENDING'), 1)
        self.assertEqual(_count_on_page(response, 'ENROLLED'), 1)

    def test_office_still_reads_the_whole_funnel(self):
        self._login(self.office_clerk, institution=self.institution_a, department='Office')
        response = self.client.get(reverse(PAGE))
        self.assertEqual(
            response.context['stage_keys'],
            ['SUBMITTED', 'OFFICE_APPROVED', 'ACCOUNT_PENDING', 'PAYMENT_APPROVED', 'ENROLLED', 'REJECTED'],
        )
        self.assertEqual(response.context['total_applications'], 6)
        self.assertFalse(response.context['payment_stages_only'])

    @requires_openpyxl
    def test_accounts_export_and_trend_keep_the_same_slice(self):
        self._login(self.accounts_clerk, institution=self.institution_a, department='Accounts')

        response = self.client.get(reverse(EXPORT))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)

        # The stage sheet carries the payment rows only — and the office words
        # do not appear anywhere in the workbook.
        counts = _funnel_sheet(workbook)
        self.assertEqual(
            list(counts),
            [LABELS['ACCOUNT_PENDING'], LABELS['PAYMENT_APPROVED'], LABELS['ENROLLED']],
        )
        text = ' '.join(
            str(cell) for sheet in workbook for row in sheet.iter_rows(values_only=True)
            for cell in row if cell is not None
        )
        # ('Submitted from'/'Submitted to' are date meta labels, not stage rows.)
        self.assertNotIn(LABELS['OFFICE_APPROVED'], text)
        self.assertNotIn(LABELS['REJECTED'], text)

        # The trend has no submission column at all, on the page or in the file.
        page = self.client.get(reverse(PAGE))
        self.assertEqual(
            [column['label'] for column in page.context['trend_columns']],
            ['Payment approved', 'Enrolled'],
        )
        self.assertNotContains(page, 'Submission date')
        header = next(
            row for row in _trend_sheet_rows(workbook)
            if row[0] == 'Bucket' and row[1] == 'Bucket start'
        )
        self.assertEqual(header, ('Bucket', 'Bucket start', 'Payment approved', 'Enrolled'))

    def test_accounts_query_param_cannot_widen_the_slice(self):
        self._login(self.accounts_clerk, institution=self.institution_a, department='Accounts')
        widened = self.client.get(reverse(PAGE), {'institution': self.institution_b.pk})
        self.assertEqual(widened.context['total_applications'], 3)
        self.assertEqual(widened.context['stage_keys'],
                         ['ACCOUNT_PENDING', 'PAYMENT_APPROVED', 'ENROLLED'])

    def test_new_sections_do_not_bypass_the_route_guards(self):
        """Hiding a menu entry is not the security — the URL itself is guarded."""
        self.client.logout()
        anonymous = self.client.get(reverse(PAGE))
        self.assertEqual(anonymous.status_code, 302)
        self.assertIn(reverse('login'), anonymous.url)

        bare = get_user_model().objects.create_user(username='of04-bare', password='password')
        InstitutionAccess.objects.create(user=bare, institution=self.institution_a, department='Office')
        self._login(bare)
        self.assertEqual(self.client.get(reverse(PAGE)).status_code, 403)
        self.assertEqual(self.client.get(reverse(EXPORT)).status_code, 403)

        exam = self._clerk('of04-exam', 'Exam')
        self._login(exam, institution=self.institution_a, department='Exam')
        for url in (reverse(PAGE), reverse(EXPORT)):
            with self.subTest(url=url):
                redirected = self.client.get(url)
                self.assertEqual(redirected.status_code, 302)
                self.assertIn(reverse('dashboard'), redirected.url)


class EmptyScopeTests(AdmissionReportFixture):
    """No data is a page of zeroes, not a 500."""

    def test_empty_scope_renders_zeroes_everywhere(self):
        Student.objects.all().delete()
        SectionCapacity.objects.all().delete()
        AdmissionApplication.objects.all().delete()
        self._login(self.admin)

        response = self.client.get(reverse(PAGE))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['capacity_rows'], [])
        self.assertEqual(response.context['trend_rows'], [])
        self.assertEqual(response.context['capacity_totals']['sections'], 0)
        self.assertContains(response, 'No seat limits or students in this scope.')
        self.assertContains(response, 'No dated events in this window.')

        export = self.client.get(reverse(EXPORT))
        self.assertEqual(export.status_code, 200)
        workbook = load_workbook(BytesIO(export.content), data_only=True)
        self.assertEqual(
            workbook.sheetnames, ['Admission Funnel', 'Capacity vs Enrolled', 'Trend'],
        )
        summary = {
            row[0]: row[1] for row in workbook['Capacity vs Enrolled'].iter_rows(values_only=True)
            if row[0]
        }
        self.assertEqual(summary['Sections listed'], 0)
        self.assertEqual(summary['Students (ACTIVE)'], 0)
        self.assertEqual(set(_funnel_sheet(workbook).values()), {0})
