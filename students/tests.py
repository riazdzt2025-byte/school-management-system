from django.test import TestCase

from io import BytesIO
from unittest import skipUnless
from datetime import date

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
try:
	from openpyxl import Workbook
except ModuleNotFoundError:
	Workbook = None

from .models import (
	AdmissionApplication, AuditLog, AttendanceRecord, Employee, EmployeeStatusLog, Exam, ExamMark, Institution, InstitutionAccess,
	MoneyReceipt, PromotionBatch, SSCRegistration, Student, StudentSubjectChoice, Subject,
)
from .forms import ExamForm, StudentForm, auto_exam_name
from .permissions import ensure_default_groups


class PermissionSetupTests(TestCase):
	def test_default_groups_are_created_with_required_permissions(self):
		ensure_default_groups()
		self.assertTrue(hasattr(Student, 'objects'))
		self.assertTrue(hasattr(AdmissionApplication, 'objects'))
		self.assertEqual(AdmissionApplication._meta.model_name, 'admissionapplication')
		self.assertTrue(hasattr(ensure_default_groups, '__call__'))


class StudentArchiveSafetyTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Archive School', classes='6,7,8')
		self.user = get_user_model().objects.create_superuser(username='archive-admin', password='password')
		self.client.force_login(self.user)
		self.student = Student.objects.create(
			institution=self.institution,
			student_id='A001',
			name='Archived Student',
			admission_class='6',
			section='A',
			admission_year=2026,
		)

	def test_bulk_delete_soft_deletes_student_and_keeps_record(self):
		response = self.client.post(reverse('bulk_delete_students'), {
			'student_ids': [self.student.pk],
			'institution': self.institution.pk,
		})
		self.assertEqual(response.url, reverse('student_list') + f'?institution={self.institution.pk}')
		self.student.refresh_from_db()
		self.assertTrue(self.student.is_archived)
		self.assertEqual(self.student.status, 'DISCONTINUED')
		self.assertTrue(Student.objects.filter(pk=self.student.pk).exists())

		list_response = self.client.get(reverse('student_list'), {'institution': self.institution.pk})
		self.assertNotContains(list_response, 'Archived Student')

	def test_restore_student_returns_it_to_active_list(self):
		self.student.status = 'TRANSFERRED'
		self.student.save(update_fields=['status'])
		self.client.post(reverse('delete_student', args=[self.student.pk]))
		self.student.refresh_from_db()
		self.assertTrue(self.student.is_archived)
		self.assertEqual(self.student.pre_archive_status, 'TRANSFERRED')

		response = self.client.post(reverse('restore_student', args=[self.student.pk]))
		self.assertEqual(response.status_code, 302)
		self.student.refresh_from_db()
		self.assertFalse(self.student.is_archived)
		# The student's status from before it was archived is restored, not a hardcoded default.
		self.assertEqual(self.student.status, 'TRANSFERRED')
		self.assertEqual(self.student.pre_archive_status, '')
		self.assertIsNotNone(self.student.restored_at)
		self.assertEqual(self.student.restored_by, self.user)

		list_response = self.client.get(reverse('student_list'), {'institution': self.institution.pk})
		self.assertContains(list_response, 'Archived Student')

	def test_archived_students_page_lists_archived_only(self):
		active_student = Student.objects.create(
			institution=self.institution, student_id='A002', name='Still Active',
			admission_class='6', section='A', admission_year=2026,
		)
		self.client.post(reverse('delete_student', args=[self.student.pk]))

		response = self.client.get(reverse('archived_students'), {'institution': self.institution.pk})
		self.assertContains(response, 'Archived Student')
		self.assertNotContains(response, 'Still Active')

	def test_bulk_restore_students(self):
		other = Student.objects.create(
			institution=self.institution, student_id='A003', name='Second Archived',
			admission_class='6', section='A', admission_year=2026,
		)
		self.client.post(reverse('bulk_delete_students'), {
			'student_ids': [self.student.pk, other.pk],
			'institution': self.institution.pk,
		})
		response = self.client.post(reverse('bulk_restore_students'), {
			'student_ids': [self.student.pk, other.pk],
			'institution': self.institution.pk,
		})
		self.assertEqual(response.status_code, 302)
		self.student.refresh_from_db()
		other.refresh_from_db()
		self.assertFalse(self.student.is_archived)
		self.assertFalse(other.is_archived)

	def test_purge_permanently_deletes_archived_student_and_related_rows(self):
		subject = Subject.objects.create(code='ENG', name='English', full_marks=100)
		exam = Exam.objects.create(
			name='Mid Term', exam_type='MID_TERM_1', institution=self.institution,
			admission_class='6', section='A', session='2026',
		)
		ExamMark.objects.create(exam=exam, student=self.student, subject=subject, marks_obtained=50)
		MoneyReceipt.objects.create(
			student=self.student, receipt_no='R-PURGE-1', purpose='Fee',
			amount=100, date=date(2026, 1, 1), created_by=self.user,
		)
		application = AdmissionApplication.objects.create(
			institution=self.institution, applicant_name='Archived Student',
			applicant_contact_no='01800000000', guardian_name='Guardian',
			guardian_contact_no='01900000000', requested_class='6',
			requested_section='A', session='2026-2027', status='ENROLLED',
			enrolled_student=self.student,
		)
		self.client.post(reverse('delete_student', args=[self.student.pk]))
		pk = self.student.pk
		response = self.client.post(reverse('purge_archived_student', args=[pk]))
		self.assertEqual(response.status_code, 302)
		self.assertFalse(Student.objects.filter(pk=pk).exists())
		self.assertEqual(ExamMark.objects.filter(exam=exam).count(), 0)
		self.assertEqual(MoneyReceipt.objects.filter(receipt_no='R-PURGE-1').count(), 0)
		application.refresh_from_db()
		self.assertIsNone(application.enrolled_student_id)
		self.assertTrue(AuditLog.objects.filter(action='student_purged', object_id=str(pk)).exists())

	def test_purge_refuses_an_active_student(self):
		response = self.client.post(reverse('purge_archived_student', args=[self.student.pk]))
		self.assertEqual(response.status_code, 404)
		self.assertTrue(Student.objects.filter(pk=self.student.pk).exists())


class PromotionAndAuditTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Promotion School', classes='6,7,8')
		self.user = get_user_model().objects.create_superuser(username='promotion-admin', password='password')
		self.client.force_login(self.user)
		self.student = Student.objects.create(institution=self.institution, student_id='P001', name='Promoted One', admission_class='6', section='A', admission_year=2026)
		self.untouched = Student.objects.create(institution=self.institution, student_id='P002', name='Promoted Two', admission_class='6', section='A', admission_year=2026)

	def test_promotion_requires_session_and_records_history(self):
		response = self.client.post(reverse('student_promotion'), {
			'from_class': '6', 'from_section': 'A', 'to_class': '7', 'to_section': 'B', 'session': '2026-2027',
		})
		self.assertRedirects(response, reverse('student_list'))
		batch = PromotionBatch.objects.get()
		self.assertEqual(batch.session, '2026-2027')
		self.assertEqual(batch.student_history.count(), 2)
		self.assertTrue(AuditLog.objects.filter(action='students_promoted', object_id=str(batch.pk)).exists())

	def test_rollback_only_restores_students_still_at_target(self):
		self.client.post(reverse('student_promotion'), {
			'from_class': '6', 'from_section': 'A', 'to_class': '7', 'to_section': 'B', 'session': '2026-2027',
		})
		self.student.admission_class = '8'
		self.student.save(update_fields=['admission_class'])
		batch = PromotionBatch.objects.get()
		response = self.client.post(reverse('rollback_student_promotion', args=[batch.pk]))
		self.assertRedirects(response, reverse('student_promotion_history'))
		self.student.refresh_from_db()
		self.untouched.refresh_from_db()
		self.assertEqual(self.student.admission_class, '8')
		self.assertEqual(self.untouched.admission_class, '6')
		self.assertEqual(batch.student_history.filter(rolled_back_at__isnull=False).count(), 1)
		self.assertTrue(AuditLog.objects.filter(action='students_promotion_rollback').exists())

	def test_promotion_rejects_same_location(self):
		response = self.client.post(reverse('student_promotion'), {
			'from_class': '6', 'from_section': 'A', 'to_class': '6', 'to_section': 'a', 'session': '2026-2027',
		})
		self.assertEqual(response.status_code, 200)
		self.assertEqual(PromotionBatch.objects.count(), 0)


class ExamWorkflowTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Test School', classes='6')
		self.student = Student.objects.create(
			institution=self.institution, student_id='S001', name='Student One',
			admission_class='6', section='A', admission_year=2026,
		)
		self.other_student = Student.objects.create(
			institution=self.institution, student_id='S002', name='Student Two',
			admission_class='7', section='A', admission_year=2026,
		)
		self.subject = Subject.objects.create(code='ENG', name='English', full_marks=100)
		from .models import SubjectRequirement
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='6', subject=self.subject,
			requirement_type='MANDATORY',
		)
		self.exam = Exam.objects.create(
			name='Mid Term', exam_type='MID_TERM_1', institution=self.institution,
			admission_class='6', section='A', session='2026',
		)
		self.user = get_user_model().objects.create_user(username='exam-user', password='password')
		exammark_content_type = ContentType.objects.get_for_model(ExamMark)
		exam_content_type = ContentType.objects.get_for_model(Exam)
		add_exammark, _ = Permission.objects.get_or_create(
			content_type=exammark_content_type,
			codename='add_exammark',
		)
		change_exam, _ = Permission.objects.get_or_create(
			content_type=exam_content_type,
			codename='change_exam',
		)
		add_exam_perm, _ = Permission.objects.get_or_create(
			content_type=exam_content_type,
			codename='add_exam',
		)
		self.user.user_permissions.add(add_exammark, change_exam, add_exam_perm)
		self.client.force_login(self.user)

	def workbook_upload(self, rows):
		workbook = Workbook()
		sheet = workbook.active
		sheet.append(['Student ID', 'Subject Code', 'Marks'])
		for row in rows:
			sheet.append(row)
		output = BytesIO()
		workbook.save(output)
		return SimpleUploadedFile(
			'marks.xlsx', output.getvalue(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		)

	def test_unpublished_result_views_redirect_to_exam_list(self):
		urls = [
			reverse('result_sheet', args=[self.exam.pk]),
			reverse('exam_result_summary', args=[self.exam.pk]),
			reverse('top_10', args=[self.exam.pk]),
			reverse('student_result_detail', args=[self.exam.pk, self.student.pk]),
			reverse('result_card', args=[self.exam.pk, self.student.pk]),
		]
		for url in urls:
			with self.subTest(url=url):
				response = self.client.get(url)
				self.assertRedirects(response, reverse('exam_list'))

	def test_publish_toggle_only_mutates_on_post(self):
		url = reverse('toggle_publish_exam', args=[self.exam.pk])
		response = self.client.get(url)
		self.assertEqual(response.status_code, 405)
		self.exam.refresh_from_db()
		self.assertFalse(self.exam.is_published)

		response = self.client.post(url)
		self.assertRedirects(response, reverse('exam_list'))
		self.exam.refresh_from_db()
		self.assertTrue(self.exam.is_published)

	def test_exam_type_uses_approved_dropdown_and_auto_names(self):
		form = ExamForm()
		self.assertNotIn('name', form.fields)
		self.assertIn(('FIRST_TERM', 'First Term'), form.fields['exam_type'].choices)
		self.assertIn(('MID_TERM_3', 'Mid Term-3'), form.fields['exam_type'].choices)

		response = self.client.get(reverse('add_exam'))
		self.assertEqual(response.status_code, 200)
		self.assertNotIn('exam_name_suggestions', response.context)

	def test_exam_form_uses_dropdowns_for_class_and_section(self):
		form = ExamForm()
		self.assertEqual(form.fields['admission_class'].widget.__class__.__name__, 'Select')
		self.assertEqual(form.fields['section'].widget.__class__.__name__, 'Select')
		self.assertIn(('A', 'A'), form.fields['section'].choices)
		self.assertIn(('J', 'J'), form.fields['section'].choices)
		self.assertNotIn(('N/A', 'N/A'), form.fields['section'].choices)

	def test_exam_clerk_cannot_delete_an_exam(self):
		permission = Permission.objects.get_or_create(codename='delete_exam', content_type=ContentType.objects.get_for_model(Exam))[0]
		self.client.force_login(self.user)
		self.user.user_permissions.add(permission)
		response = self.client.post(reverse('delete_exam', args=[self.exam.pk]))
		self.assertRedirects(response, reverse('exam_list'))
		self.assertTrue(Exam.objects.filter(pk=self.exam.pk).exists())
		list_page = self.client.get(reverse('exam_list'))
		self.assertNotContains(list_page, 'Delete')

	def test_admin_can_delete_a_mistaken_exam_and_its_marks(self):
		ExamMark.objects.create(
			exam=self.exam, student=self.student, subject=self.subject, marks_obtained=40,
		)
		admin = get_user_model().objects.create_superuser(username='exam-admin', password='password')
		self.client.force_login(admin)
		list_page = self.client.get(reverse('exam_list'))
		self.assertContains(list_page, 'Delete')
		response = self.client.post(reverse('delete_exam', args=[self.exam.pk]))
		self.assertRedirects(response, reverse('exam_list'))
		self.assertFalse(Exam.objects.filter(pk=self.exam.pk).exists())
		self.assertEqual(ExamMark.objects.filter(student=self.student).count(), 0)
		self.assertTrue(AuditLog.objects.filter(action='exam_deleted').exists())

	@skipUnless(Workbook, 'openpyxl is required for Excel import tests')
	def test_import_exam_marks_success(self):
		response = self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.subject.pk),
				'excel_file': self.workbook_upload([['S001', 'ENG', 87.5]]),
			},
		)
		self.assertRedirects(response, reverse('exam_list'))
		mark = ExamMark.objects.get(exam=self.exam, student=self.student, subject=self.subject)
		self.assertEqual(str(mark.marks_obtained), '87.50')

	@skipUnless(Workbook, 'openpyxl is required for Excel import tests')
	def test_invalid_row_skips_students_outside_this_exam(self):
		response = self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.subject.pk),
				'excel_file': self.workbook_upload([
					['S001', 'ENG', 75],
					['S002', 'ENG', 80],
				]),
			},
		)
		self.assertRedirects(response, reverse('exam_list'))
		self.assertEqual(ExamMark.objects.filter(exam=self.exam).count(), 1)
		self.assertEqual(str(ExamMark.objects.get().marks_obtained), '75.00')

	@skipUnless(Workbook, 'openpyxl is required for Excel import tests')
	def test_import_matches_roll_and_name_when_the_student_id_changed(self):
		self.student.roll_no = 117
		self.student.save(update_fields=['roll_no'])
		Student.objects.create(
			institution=self.institution, student_id='202609137', name='Old Khalid Record',
			admission_class='8', section='A', admission_year=2026, roll_no=117,
		)
		workbook = Workbook()
		sheet = workbook.active
		sheet.append(['Roll', 'ID', 'Name', 'Marks'])
		sheet.append([117, '202609137', 'Student One', 81])
		output = BytesIO()
		workbook.save(output)
		response = self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.subject.pk),
				'excel_file': SimpleUploadedFile(
					'bangla.xlsx', output.getvalue(),
					content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
				),
			},
		)
		self.assertRedirects(response, reverse('exam_list'))
		mark = ExamMark.objects.get(exam=self.exam, student=self.student, subject=self.subject)
		self.assertEqual(str(mark.marks_obtained), '81.00')

	def test_import_page_asks_for_a_subject_first(self):
		response = self.client.get(reverse('import_exam_marks', args=[self.exam.pk]))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'one subject')
		self.assertContains(response, 'name="subject"')
		self.assertNotContains(response, 'Upload and Import')

	@skipUnless(Workbook, 'openpyxl is required for Excel import tests')
	def test_import_skips_students_who_are_no_longer_on_the_roll(self):
		response = self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.subject.pk),
				'excel_file': self.workbook_upload([
					['S001', 'ENG', 70],
					['GONE99', 'ENG', 80],
				]),
			},
		)
		self.assertRedirects(response, reverse('exam_list'))
		self.assertEqual(ExamMark.objects.filter(exam=self.exam).count(), 1)
		self.assertEqual(str(ExamMark.objects.get().marks_obtained), '70.00')



class StudentProfileAndBulkUpdateTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Profile Campus', classes='6,7,8,9,10')
		self.user = get_user_model().objects.create_superuser(username='profile-admin', password='password')
		self.client.force_login(self.user)
		self.student = Student.objects.create(
			institution=self.institution,
			student_id='P101',
			name='Profile Student',
			admission_class='9',
			section='A',
			group='SCI',
			admission_year=2026,
		)

	def test_student_form_includes_photo_field(self):
		self.assertIn('photo', StudentForm().fields)

	def test_bulk_update_select_has_class_and_group_controls(self):
		response = self.client.post(reverse('bulk_update_select'), {
			'student_ids': [self.student.pk],
			'institution': self.institution.pk,
			'admission_class': '9',
			'section': 'A',
		})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'name="new_group"')
		self.assertContains(response, 'name="new_class"')
		self.assertContains(response, 'name="new_section"')


class AttendanceTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Attendance Campus', classes='6,7')
		self.student = Student.objects.create(
			institution=self.institution, student_id='A001', name='Daily Student',
			admission_class='6', section='A', admission_year=2026,
		)
		self.user = get_user_model().objects.create_user(username='attendance-user', password='password')
		self.client.force_login(self.user)

	def test_student_attendance_record_can_be_created(self):
		record = AttendanceRecord.objects.create(
			institution=self.institution,
			student=self.student,
			date='2026-09-01',
			status='P',
			remarks='Present',
		)
		self.assertEqual(record.status, 'P')
		self.assertEqual(record.student.name, 'Daily Student')

	def test_attendance_summary_page_loads(self):
		AttendanceRecord.objects.create(
			institution=self.institution,
			student=self.student,
			date='2026-09-01',
			status='P',
		)
		response = self.client.get(reverse('attendance_report'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Attendance Report')

	def test_mark_attendance_selection_page_loads(self):
		response = self.client.get(reverse('mark_attendance'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Mark Attendance')

	def test_mark_attendance_bulk_page_renders_students(self):
		InstitutionAccess.objects.create(
			user=self.user, institution=self.institution, department='Office', is_active=True
		)
		from datetime import date
		today = date.today().isoformat()
		response = self.client.get(
			reverse('mark_attendance_bulk', kwargs={
				'date_str': today, 'admission_class': '6', 'section': 'A', 'mark_type': 'STUDENT'
			})
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, self.student.name)
		self.assertContains(response, 'Status')

	def test_mark_attendance_bulk_updates_records(self):
		InstitutionAccess.objects.create(
			user=self.user, institution=self.institution, department='Office', is_active=True
		)
		from datetime import date
		today = date.today().isoformat()
		response = self.client.post(
			reverse('mark_attendance_bulk', kwargs={
				'date_str': today, 'admission_class': '6', 'section': 'A', 'mark_type': 'STUDENT'
			}),
			{f'status_{self.student.id}': 'P', f'remarks_{self.student.id}': 'Present in class'}
		)
		self.assertRedirects(response, reverse('attendance_report'))
		record = AttendanceRecord.objects.get(student=self.student, date=today)
		self.assertEqual(record.status, 'P')
		self.assertEqual(record.remarks, 'Present in class')

	def test_attendance_summary_page_loads_and_shows_statistics(self):
		"""Test that attendance summary page loads with statistics."""
		from datetime import date, timedelta
		InstitutionAccess.objects.create(
			user=self.user, institution=self.institution, department='Office', is_active=True
		)
		today = date.today()
		yesterday = today - timedelta(days=1)
		
		# Create multiple attendance records on different dates
		AttendanceRecord.objects.create(
			student=self.student, date=today, status='P', 
			institution=self.institution, created_by=self.user
		)
		AttendanceRecord.objects.create(
			student=self.student, date=yesterday, status='A',
			institution=self.institution, created_by=self.user
		)
		
		response = self.client.get(reverse('attendance_summary'))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Attendance Summary')
		self.assertContains(response, self.student.name)


class StudentDetailPageTests(TestCase):
	def setUp(self):
		self.user = get_user_model().objects.create_superuser(username='admin', password='password')
		self.client.force_login(self.user)
		self.institution = Institution.objects.create(name='Test School', classes='6,7,8')
		self.student = Student.objects.create(
			name='Detail Test Student', student_id='D001', admission_class='6', section='A',
			gender='M', institution=self.institution, created_by=self.user
		)

	def test_student_detail_page_loads(self):
		"""Test that student detail page loads successfully."""
		response = self.client.get(reverse('student_detail', args=[self.student.pk]))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Detail Test Student')
		self.assertContains(response, 'D001')
		self.assertContains(response, 'Subjects & Curriculum')

	def test_student_detail_shows_attendance_records(self):
		"""Test that student detail page displays attendance records."""
		from datetime import date
		AttendanceRecord.objects.create(
			student=self.student, date=date.today(), status='P',
			institution=self.institution, created_by=self.user
		)
		
		response = self.client.get(reverse('student_detail', args=[self.student.pk]))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Attendance')
		self.assertContains(response, 'Present')

	def test_student_detail_shows_exam_results(self):
		"""Test that student detail page displays exam results."""
		exam = Exam.objects.create(
			name='Quarterly', exam_type='MID_TERM', institution=self.institution,
			admission_class='6', section='A', session='2025-2026'
		)
		subject = Subject.objects.create(name='Math', code='MATH001', created_by=self.user)
		ExamMark.objects.create(
			student=self.student, exam=exam, subject=subject, marks_obtained=85
		)
		
		response = self.client.get(reverse('student_detail', args=[self.student.pk]))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Exam Results')
		self.assertContains(response, 'Math')


class EmployeeDetailPageTests(TestCase):
	def setUp(self):
		self.user = get_user_model().objects.create_superuser(username='admin', password='password')
		self.client.force_login(self.user)
		self.institution = Institution.objects.create(name='Test School', classes='6,7,8')
		self.employee = Employee.objects.create(
			name='John Doe', designation='Teacher', department='Science',
			institution=self.institution, join_date=date(2020, 1, 15), contact_no='01700000000'
		)

	def test_employee_detail_page_loads(self):
		"""Test that employee detail page loads successfully."""
		response = self.client.get(reverse('employee_detail', args=[self.employee.pk]))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'John Doe')
		self.assertContains(response, 'Teacher')
		self.assertContains(response, 'Overview')

	def test_employee_detail_shows_status_history(self):
		"""Test that employee detail page displays status history."""
		EmployeeStatusLog.objects.create(
			employee=self.employee, old_status='ACTIVE', new_status='OSD',
			reason='On Special Duty', changed_by=self.user
		)
		
		response = self.client.get(reverse('employee_detail', args=[self.employee.pk]))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Status History')
		self.assertContains(response, 'On Special Duty')

	def test_employee_detail_shows_attendance_records(self):
		"""Test that employee detail page displays attendance records."""
		from datetime import date
		AttendanceRecord.objects.create(
			employee=self.employee, date=date.today(), status='P',
			institution=self.institution, created_by=self.user
		)
		
		response = self.client.get(reverse('employee_detail', args=[self.employee.pk]))
		
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Attendance')
		self.assertContains(response, 'Present')


class InstitutionAwareLoginTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Trust Campus', classes='6,7,8')
		self.other_institution = Institution.objects.create(name='City Campus', classes='6,7,8')
		self.user = get_user_model().objects.create_user(username='office_user', password='password')
		InstitutionAccess.objects.create(user=self.user, institution=self.institution, department='Office')
		InstitutionAccess.objects.create(user=self.user, institution=self.other_institution, department='Exam')

	def test_login_page_lists_only_accessible_institutions(self):
		response = self.client.get(reverse('login'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Trust Campus')
		self.assertContains(response, 'City Campus')
		self.assertContains(response, 'Office')
		self.assertContains(response, 'Exam')

	def test_selected_institution_and_department_are_saved_on_login(self):
		response = self.client.post(reverse('login'), {
			'username': 'office_user',
			'password': 'password',
			'institution_id': str(self.institution.pk),
			'department': 'Office',
		})
		self.assertEqual(response.status_code, 302)
		self.assertEqual(self.client.session.get('selected_institution_id'), str(self.institution.pk))
		self.assertEqual(self.client.session.get('selected_department'), 'Office')


class UIConsistencyTests(TestCase):
	def setUp(self):
		self.user = get_user_model().objects.create_user(username='ui_user', password='password')
		self.client.force_login(self.user)

	def test_dashboard_excludes_quick_access_section_and_duplicate_logout(self):
		response = self.client.get(reverse('dashboard'))
		self.assertEqual(response.status_code, 200)
		self.assertNotContains(response, 'Quick Search')
		self.assertNotContains(response, 'Quick Access')
		self.assertEqual(response.content.decode('utf-8').count('Logout'), 1)

	def test_selected_institution_filters_student_list(self):
		inst_a = Institution.objects.create(name='Alpha Campus', classes='6,7')
		inst_b = Institution.objects.create(name='Beta Campus', classes='6,7')
		InstitutionAccess.objects.create(user=self.user, institution=inst_a, department='Office')
		InstitutionAccess.objects.create(user=self.user, institution=inst_b, department='Exam')
		Student.objects.create(institution=inst_a, student_id='A001', name='Alpha Student', admission_class='6', section='A', admission_year=2026)
		Student.objects.create(institution=inst_b, student_id='B001', name='Beta Student', admission_class='6', section='A', admission_year=2026)
		session = self.client.session
		session['selected_institution_id'] = str(inst_a.pk)
		session['selected_department'] = 'Office'
		session.save()

		response = self.client.get(reverse('student_list'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Alpha Student')
		self.assertNotContains(response, 'Beta Student')


class AdmissionApplicationWorkflowTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Admission School', classes='6,7')
		self.user = get_user_model().objects.create_superuser(
			username='admission-admin', password='password', email='admin@example.com'
		)
		self.client.force_login(self.user)
		self.application = AdmissionApplication.objects.create(
			institution=self.institution, applicant_name='Applicant One',
			applicant_contact_no='01800000000', guardian_name='Guardian One',
			guardian_contact_no='01900000000', requested_class='6',
			requested_section='A', session='2026-2027',
		)

	def post_transition(self, name, data=None):
		return self.client.post(reverse(name, args=[self.application.pk]), data or {})

	def test_office_transition_requires_post_and_handoff(self):
		url = reverse('office_approve_application', args=[self.application.pk])
		self.assertEqual(self.client.get(url).status_code, 405)
		self.post_transition('office_approve_application', {'remarks': 'Verified'})
		self.application.refresh_from_db()
		self.assertEqual(self.application.status, 'OFFICE_APPROVED')
		self.post_transition('office_handoff_application')
		self.application.refresh_from_db()
		self.assertEqual(self.application.status, 'ACCOUNT_PENDING')

	def test_invalid_transition_does_not_mutate(self):
		self.post_transition('office_handoff_application')
		self.application.refresh_from_db()
		self.assertEqual(self.application.status, 'SUBMITTED')
		self.assertEqual(Student.objects.count(), 0)

	def test_accounts_queue_filters_by_requested_class(self):
		self.application.status = 'ACCOUNT_PENDING'
		self.application.save(update_fields=['status'])
		other = AdmissionApplication.objects.create(
			institution=self.institution, applicant_name='Applicant Two',
			applicant_contact_no='01800000001', guardian_name='Guardian Two',
			guardian_contact_no='01900000001', requested_class='7', session='2026-2027',
			status='ACCOUNT_PENDING',
		)
		response = self.client.get(reverse('accounts_admission_queue'), {'admission_class': '6'})
		self.assertContains(response, self.application.application_number)
		self.assertNotContains(response, other.application_number)

	def test_payment_approval_creates_one_student_and_receipt(self):
		self.application.status = 'ACCOUNT_PENDING'
		self.application.save(update_fields=['status'])
		data = {
			'payment_amount': '1500.00', 'payment_date': '2026-08-27',
			'payment_purpose': 'Admission Fee', 'account_remarks': 'Paid',
		}
		response = self.post_transition('accounts_approve_payment', data)
		self.assertRedirects(response, reverse('admission_application_detail', args=[self.application.pk]))
		self.application.refresh_from_db()
		self.assertEqual(self.application.status, 'ENROLLED')
		self.assertIsNotNone(self.application.enrolled_student_id)
		self.assertEqual(Student.objects.filter(admission_application=self.application).count(), 1)
		self.assertEqual(MoneyReceipt.objects.filter(student=self.application.enrolled_student).count(), 1)
		self.assertEqual(MoneyReceipt.objects.filter(student=self.application.enrolled_student).first().created_by, self.user)
		self.assertEqual(self.post_transition('accounts_approve_payment', data).status_code, 302)
		self.assertEqual(Student.objects.filter(admission_application=self.application).count(), 1)
		self.assertEqual(MoneyReceipt.objects.filter(student=self.application.enrolled_student).count(), 1)


class DepartmentAccessControlTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Access Control School', classes='6,7')
		self.office_user = get_user_model().objects.create_user(username='office_user', password='password')
		self.accounts_user = get_user_model().objects.create_user(username='accounts_user', password='password')
		InstitutionAccess.objects.create(user=self.office_user, institution=self.institution, department='Office')
		InstitutionAccess.objects.create(user=self.accounts_user, institution=self.institution, department='Accounts')

	def test_office_and_accounts_users_can_view_admission_list(self):
		# Both Office and Accounts users can view admission list.
		admission_app_ct = ContentType.objects.get_for_model(AdmissionApplication)
		perm = Permission.objects.get(content_type=admission_app_ct, codename='view_admissionapplication')
		for user in [self.office_user, self.accounts_user]:
			user.user_permissions.add(perm)
		
		for user, dept in [(self.office_user, 'Office'), (self.accounts_user, 'Accounts')]:
			self.client.force_login(user)
			session = self.client.session
			session['selected_institution_id'] = str(self.institution.pk)
			session['selected_department'] = dept
			session.save()
			response = self.client.get(reverse('admission_application_list'))
			self.assertEqual(response.status_code, 200)

	def test_institution_filtering_works_in_admission_list(self):
		# Admission applications are filtered by selected institution.
		inst1 = Institution.objects.create(name='Inst1', classes='6')
		self.office_user.user_permissions.add(
			Permission.objects.get(content_type=ContentType.objects.get_for_model(AdmissionApplication), codename='view_admissionapplication')
		)
		InstitutionAccess.objects.create(user=self.office_user, institution=inst1, department='Office')
		
		app1 = AdmissionApplication.objects.create(
			institution=inst1, applicant_name='App1', applicant_contact_no='01800000000',
			guardian_name='Guard1', guardian_contact_no='01900000000', requested_class='6', session='2026-2027'
		)
		
		self.client.force_login(self.office_user)
		session = self.client.session
		session['selected_institution_id'] = str(inst1.pk)
		session['selected_department'] = 'Office'
		session.save()
		
		response = self.client.get(reverse('admission_application_list'))
		self.assertEqual(response.status_code, 200)
		self.assertIn(app1.application_number, str(response.content))

# Create your tests here.



class MarksPartsAndPassRulesTests(TestCase):
	"""Practical / Weekly Test entry, the part pass rule, and blank-is-absent."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Marks School', classes='9')
		self.user = get_user_model().objects.create_superuser(username='marks-admin', password='password')
		self.client.force_login(self.user)
		self.physics = Subject.objects.create(code='PHY', name='Physics', full_marks=100)
		from .models import SubjectRequirement
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', group='SCI',
			subject=self.physics, requirement_type='MANDATORY',
		)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='9', section='', group='SCI',
			session='2026',
		)
		self.student = Student.objects.create(
			institution=self.institution, student_id='M001', name='Marked Student',
			admission_class='9', section='A', group='SCI', roll_no=1, admission_year=2026,
		)
		self.absent_student = Student.objects.create(
			institution=self.institution, student_id='M002', name='Absent Student',
			admission_class='9', section='A', group='SCI', roll_no=2, admission_year=2026,
		)

	def configure(self, **overrides):
		fields = {
			'institution': self.institution, 'admission_class': '9',
			'subject': self.physics, 'exam_type': 'SECOND_TERM',
			'full_marks': 100, 'cq_marks': 75, 'mcq_marks': 0, 'practical_marks': 25,
		}
		fields.update(overrides)
		from .models import SubjectMarkSetting
		return SubjectMarkSetting.objects.create(**fields)

	def post_marks(self, payload):
		return self.client.post(
			reverse('enter_marks', args=[self.exam.pk, self.physics.pk]), payload
		)

	def test_parts_are_added_into_the_total(self):
		self.configure()
		self.post_marks({
			f'cq_{self.student.pk}': '60',
			f'mcq_{self.student.pk}': '',
			f'practical_{self.student.pk}': '20',
		})
		mark = ExamMark.objects.get(exam=self.exam, student=self.student, subject=self.physics)
		self.assertEqual(str(mark.marks_obtained), '80.00')
		self.assertEqual(str(mark.cq_obtained), '60.00')
		self.assertEqual(str(mark.practical_obtained), '20.00')
		self.assertIsNone(mark.weekly_test_obtained)

	def test_a_student_entered_nowhere_has_no_row(self):
		"""Blank everywhere means absent: no row, so the result shows a dash."""
		self.configure()
		self.post_marks({
			f'cq_{self.student.pk}': '60', f'practical_{self.student.pk}': '20',
			f'cq_{self.absent_student.pk}': '', f'practical_{self.absent_student.pk}': '',
		})
		self.assertFalse(ExamMark.objects.filter(student=self.absent_student).exists())

	def test_mark_above_a_part_max_is_not_saved(self):
		self.configure()
		response = self.post_marks({
			f'cq_{self.student.pk}': '90', f'practical_{self.student.pk}': '20',
		})
		self.assertEqual(response.status_code, 302)
		self.assertFalse(ExamMark.objects.filter(student=self.student).exists())
		followed = self.client.get(reverse('exam_list'))
		self.assertContains(followed, 'must be between 0 and 75')

	def test_each_part_must_pass_fails_the_subject_on_one_part(self):
		setting = self.configure(require_all_parts_pass=True, pass_percentage=40)
		# 80/100 overall passes the subject total, but 5/25 fails the practical
		# (pass mark for that part is 10).
		ExamMark.objects.create(
			exam=self.exam, student=self.student, subject=self.physics,
			marks_obtained=80, cq_obtained=75, practical_obtained=5,
		)
		from .result_utils import build_exam_results, get_subject_marks
		result = compute_first_result(self.exam)
		self.assertEqual(result['status'], 'Fail')
		self.assertEqual(str(result['gpa']), '0.00')
		subject_result = result['subject_results'][0]
		self.assertEqual(subject_result['grade'], 'F')
		self.assertEqual(subject_result['failed_parts'], ['Practical'])
		self.assertEqual(setting.pass_marks, 40.0)
		self.assertEqual(get_subject_marks(self.exam, self.physics).part_pass_marks(25), 10.0)

	def test_without_the_tick_only_the_total_matters(self):
		self.configure(require_all_parts_pass=False)
		ExamMark.objects.create(
			exam=self.exam, student=self.student, subject=self.physics,
			marks_obtained=80, cq_obtained=75, practical_obtained=5,
		)
		result = compute_first_result(self.exam)
		self.assertEqual(result['status'], 'Pass')

	def test_configured_but_blank_part_is_a_failed_part(self):
		self.configure(require_all_parts_pass=True)
		ExamMark.objects.create(
			exam=self.exam, student=self.student, subject=self.physics,
			marks_obtained=99, cq_obtained=74, mcq_obtained=25, practical_obtained=None,
		)
		result = compute_first_result(self.exam)
		self.assertEqual(result['status'], 'Fail')
		self.assertEqual(result['subject_results'][0]['failed_parts'], ['Practical'])

	def test_parts_total_mismatch_is_reported_on_the_entry_page(self):
		self.configure(practical_marks=None, cq_marks=70, mcq_marks=20)
		response = self.client.get(reverse('enter_marks', args=[self.exam.pk, self.physics.pk]))
		self.assertContains(response, 'add up to 90')

	@skipUnless(Workbook, 'openpyxl is required for Excel import tests')
	def test_import_writes_cq_mcq_pt_parts(self):
		self.configure()
		workbook = Workbook()
		sheet = workbook.active
		sheet.title = '9SC Physics'
		sheet.append(['Roll', 'ID', 'Name', 'CQ', 'MCQ', 'PT'])
		sheet.append([1, 'M001', 'Marked Student', 60, None, 20])
		output = BytesIO()
		workbook.save(output)
		response = self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.physics.pk),
				'excel_file': SimpleUploadedFile(
					'9SC Physics.xlsx', output.getvalue(),
					content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
				),
			},
		)
		self.assertRedirects(response, reverse('exam_list'))
		mark = ExamMark.objects.get(exam=self.exam, student=self.student, subject=self.physics)
		self.assertEqual(str(mark.marks_obtained), '80.00')
		self.assertEqual(str(mark.cq_obtained), '60.00')
		self.assertEqual(str(mark.practical_obtained), '20.00')
		self.assertIsNone(mark.mcq_obtained)

	def test_weekly_test_is_an_enterable_part(self):
		self.configure(cq_marks=60, mcq_marks=20, weekly_test_marks=20)
		self.post_marks({
			f'cq_{self.student.pk}': '50', f'mcq_{self.student.pk}': '15',
			f'weekly_test_{self.student.pk}': '18',
		})
		mark = ExamMark.objects.get(exam=self.exam, student=self.student)
		self.assertEqual(str(mark.marks_obtained), '83.00')
		self.assertEqual(str(mark.weekly_test_obtained), '18.00')


def compute_first_result(exam):
	"""Result row for the first student, shared by the tests above."""
	from .result_utils import build_exam_results
	_, results = build_exam_results(exam)
	return results[0]


class ExamScopeConsistencyTests(TestCase):
	"""Marks entry, import, seat plan and results must list the same students."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Scope School', classes='9')
		self.user = get_user_model().objects.create_superuser(username='scope-admin', password='password')
		self.client.force_login(self.user)
		self.physics = Subject.objects.create(code='PHY2', name='Physics', full_marks=100)
		self.accounting = Subject.objects.create(code='ACC2', name='Accounting', full_marks=100)
		from .models import SubjectRequirement
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', group='SCI',
			subject=self.physics, requirement_type='MANDATORY',
		)
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', group='BUS',
			subject=self.accounting, requirement_type='MANDATORY',
		)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='9', group='SCI', session='2026',
		)
		# Zero-padded class: the exam module must still match these students.
		self.science_student = Student.objects.create(
			institution=self.institution, student_id='S010', name='Science Kid',
			admission_class='09', section='A', group='SCI', roll_no=1, admission_year=2026,
		)
		self.business_student = Student.objects.create(
			institution=self.institution, student_id='B010', name='Business Kid',
			admission_class='9', section='A', group='BUS', roll_no=1, admission_year=2026,
		)
		self.archived_student = Student.objects.create(
			institution=self.institution, student_id='S011', name='Gone Student',
			admission_class='9', section='A', group='SCI', roll_no=2,
			admission_year=2026, is_archived=True,
		)

	def test_get_exam_students_scopes_by_class_group_and_archive(self):
		from .result_utils import get_exam_students
		self.assertEqual(list(get_exam_students(self.exam)), [self.science_student])

	def test_enter_marks_and_seat_plan_agree(self):
		marks_page = self.client.get(reverse('enter_marks', args=[self.exam.pk, self.physics.pk]))
		self.assertContains(marks_page, 'Science Kid')
		self.assertNotContains(marks_page, 'Business Kid')
		self.assertNotContains(marks_page, 'Gone Student')

		seat_page = self.client.get(reverse('seat_plan_list', args=[self.exam.pk]))
		self.assertContains(seat_page, '1')
		self.assertEqual(seat_page.context['total_students'], 1)

	def test_import_rejects_other_group_student_and_unassigned_subject(self):
		try:
			from openpyxl import Workbook
		except ModuleNotFoundError:
			self.skipTest('openpyxl is required for Excel import tests')
		from io import BytesIO
		from django.core.files.uploadedfile import SimpleUploadedFile

		workbook = Workbook()
		sheet = workbook.active
		sheet.append(['Student ID', 'Subject Code', 'Marks'])
		sheet.append(['B010', 'PHY2', 55])   # Business student in a Science exam
		output = BytesIO()
		workbook.save(output)
		response = self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.physics.pk),
				'excel_file': SimpleUploadedFile(
					'marks.xlsx', output.getvalue(),
					content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
				),
			},
		)
		self.assertRedirects(response, reverse('exam_list'))
		self.assertEqual(ExamMark.objects.count(), 0)

	def test_import_skips_rows_without_a_mark(self):
		try:
			from openpyxl import Workbook
		except ModuleNotFoundError:
			self.skipTest('openpyxl is required for Excel import tests')
		from io import BytesIO
		from django.core.files.uploadedfile import SimpleUploadedFile

		workbook = Workbook()
		sheet = workbook.active
		sheet.append(['Student ID', 'Subject Code', 'Marks'])
		sheet.append(['S010', 'PHY2', 71])
		sheet.append(['S010', 'ACC2', None])   # blank = not entered, not an error
		output = BytesIO()
		workbook.save(output)
		self.client.post(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{
				'subject': str(self.physics.pk),
				'excel_file': SimpleUploadedFile(
					'marks.xlsx', output.getvalue(),
					content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
				),
			},
		)
		self.assertEqual(ExamMark.objects.count(), 1)
		self.assertEqual(str(ExamMark.objects.get().marks_obtained), '71.00')

	def test_import_uses_the_exam_full_marks_not_the_subject_default(self):
		from .models import SubjectMarkSetting
		SubjectMarkSetting.objects.create(
			institution=self.institution, admission_class='9', subject=self.physics,
			exam_type='SECOND_TERM', full_marks=50,
		)
		self.assertTrue(self.physics.full_marks > 50)
		# (guard: the 71 mark used above would be rejected against a 50-mark exam)
		from .result_utils import get_subject_marks
		self.assertEqual(get_subject_marks(self.exam, self.physics).full_marks, 50)

	def test_result_sheet_excludes_other_group_subjects_and_warns(self):
		ExamMark.objects.create(
			exam=self.exam, student=self.science_student, subject=self.accounting,
			marks_obtained=80,
		)
		self.exam.is_published = True
		self.exam.save()
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		self.assertContains(response, 'Physics')
		content = response.content.decode()
		self.assertNotIn('<th>Accounting', content)
		self.assertContains(response, 'not assigned to its class/group')

	def test_zero_gpa_students_are_not_ranked_and_do_not_crash_top_10(self):
		self.exam.is_published = True
		self.exam.save()
		ExamMark.objects.create(
			exam=self.exam, student=self.science_student, subject=self.physics, marks_obtained=80,
		)
		response = self.client.get(reverse('top_10', args=[self.exam.pk]))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Science Kid')


class MarksImportTemplateDownloadTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='Template School', classes='9')
		self.user = get_user_model().objects.create_superuser(username='template-admin', password='password')
		self.client.force_login(self.user)
		self.subject = Subject.objects.create(code='BNG', name='Bangla', full_marks=100)
		from .models import SubjectRequirement
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', group='SCI',
			subject=self.subject, requirement_type='MANDATORY',
		)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='9', group='SCI', session='2026',
		)
		Student.objects.create(
			institution=self.institution, student_id='T001', name='Template Kid',
			admission_class='9', section='A', group='SCI', roll_no=1, admission_year=2026,
		)

	def test_workbook_matches_what_the_importer_accepts(self):
		try:
			from openpyxl import load_workbook
		except ModuleNotFoundError:
			self.skipTest('openpyxl is required for Excel export tests')
		response = self.client.get(
			reverse('download_marks_import_template', args=[self.exam.pk]),
			{'subject': self.subject.pk},
		)
		self.assertEqual(response.status_code, 200)
		self.assertIn('spreadsheetml.sheet', response['Content-Type'])
		import io
		workbook = load_workbook(io.BytesIO(response.content), read_only=True)
		self.assertEqual(workbook.sheetnames[0], '9SC Bangla')
		sheet = workbook['9SC Bangla']
		rows = list(sheet.iter_rows(values_only=True))
		self.assertEqual(rows[0], ('Roll', 'ID', 'Name', 'Marks'))
		self.assertEqual(rows[1], (1, 'T001', 'Template Kid', None))
		self.assertIn('How to fill', workbook.sheetnames)

	def test_template_picks_up_students_who_join_or_leave(self):
		try:
			from openpyxl import load_workbook
		except ModuleNotFoundError:
			self.skipTest('openpyxl is required for Excel export tests')
		import io
		Student.objects.create(
			institution=self.institution, student_id='T002', name='New Arrival',
			admission_class='9', section='A', group='SCI', roll_no=2, admission_year=2026,
		)
		url = reverse('download_marks_import_template', args=[self.exam.pk])
		first = load_workbook(io.BytesIO(
			self.client.get(url, {'subject': self.subject.pk}).content
		), read_only=True)
		names = [row[2] for row in first[first.sheetnames[0]].iter_rows(min_row=2, values_only=True)]
		self.assertIn('Template Kid', names)
		self.assertIn('New Arrival', names)

		Student.objects.filter(student_id='T001').update(is_archived=True)
		second = load_workbook(io.BytesIO(
			self.client.get(url, {'subject': self.subject.pk}).content
		), read_only=True)
		names = [row[2] for row in second[second.sheetnames[0]].iter_rows(min_row=2, values_only=True)]
		self.assertEqual(names, ['New Arrival'])
		self.assertEqual(
			self.client.get(url, {'subject': self.subject.pk})['Cache-Control'],
			'no-store, no-cache, must-revalidate, max-age=0',
		)

	def test_import_page_shows_the_live_class_roll(self):
		response = self.client.get(
			reverse('import_exam_marks', args=[self.exam.pk]),
			{'subject': self.subject.pk},
		)
		self.assertContains(response, 'Template Kid')
		self.assertContains(response, 'T001')
		self.assertContains(response, 'current')


class ExamNamingAndLayoutTests(TestCase):
	def setUp(self):
		self.user = get_user_model().objects.create_superuser(username='layout-admin', password='password')
		self.client.force_login(self.user)
		self.institution = Institution.objects.create(name='Layout School', classes='9')
		self.exam = Exam.objects.create(
			name='Placeholder', exam_type='SECOND_TERM', institution=self.institution,
			admission_class='9', session='2026-2027',
		)

	def test_auto_exam_name_uses_the_last_year_in_the_session(self):
		self.assertEqual(auto_exam_name('SECOND_TERM', '2026-2027'), 'Second Term Examination-2027')
		self.assertEqual(auto_exam_name('MID_TERM_1', '2026'), 'Mid Term-1 Examination-2026')
		self.assertEqual(auto_exam_name('MODEL_TEST_3', ''), 'Model Test-3 Examination')

	def test_saving_the_exam_form_regenerates_the_name(self):
		from .forms import ExamForm
		form = ExamForm(instance=self.exam, data={
			'institution': self.institution.pk, 'admission_class': '9', 'section': '',
			'group': '', 'exam_type': 'FINAL_TERM', 'session': '2026-2027',
		})
		self.assertTrue(form.is_valid(), form.errors)
		form.save()
		self.exam.refresh_from_db()
		self.assertEqual(self.exam.name, 'Final Term Examination-2027')

	def test_exam_form_offers_section_as_a_dropdown(self):
		from .forms import ExamForm
		form = ExamForm()
		self.assertEqual(form.fields['section'].widget.__class__.__name__, 'Select')
		self.assertNotIn('name', form.fields)

	def test_audit_log_and_import_students_render_in_the_app_layout(self):
		for name, url_name in (('Audit Log', 'audit_log_list'), ('Import Students', 'import_students')):
			with self.subTest(page=name):
				response = self.client.get(reverse(url_name))
				self.assertEqual(response.status_code, 200)
				self.assertContains(response, 'Principal Kazi Faruky School And College')
				self.assertContains(response, 'sidebar')

	def test_no_auto_field_warnings_remain(self):
		from django.core.checks import run_checks
		warnings = [issue for issue in run_checks() if issue.id.startswith('models.W042')]
		self.assertEqual(warnings, [])


class SSCGroupAlignmentTests(TestCase):
	def setUp(self):
		self.institution = Institution.objects.create(name='SSC School', classes='9,10')
		self.user = get_user_model().objects.create_superuser(username='ssc-admin', password='password')
		self.client.force_login(self.user)
		self.student = Student.objects.create(
			institution=self.institution, student_id='SSC01', name='SSC Kid',
			admission_class='10', section='A', group='SCI', roll_no=1, admission_year=2026,
		)

	def test_ssc_group_codes_are_the_student_group_codes(self):
		from .models import SSCRegistration
		student_codes = {code for code, _label in Student.GROUP_CHOICES}
		self.assertTrue(SSCRegistration.GROUP_CHOICES)
		for code, _label in SSCRegistration.GROUP_CHOICES:
			self.assertIn(code, student_codes)

	def test_form_refuses_a_group_that_disagrees_with_the_student(self):
		from .forms import SSCRegistrationForm
		from .models import SSCRegistration
		form = SSCRegistrationForm(student=self.student, data={
			'registration_number': 'R-1', 'session': '2025-2026', 'group': 'BUS',
			'board': 'DHAKA', 'subjects': '', 'roll_number': '', 'center': '',
		})
		self.assertFalse(form.is_valid())
		self.assertIn('group', form.errors)
		self.assertIn('Science', form.errors['group'][0])

	def test_form_accepts_a_matching_group_and_prefills_it(self):
		from .forms import SSCRegistrationForm
		form = SSCRegistrationForm(student=self.student)
		self.assertEqual(form.fields['group'].initial, 'SCI')
		form = SSCRegistrationForm(student=self.student, data={
			'registration_number': 'R-2', 'session': '2025-2026', 'group': 'SCI',
			'board': 'DHAKA', 'subjects': 'Bangla', 'roll_number': '', 'center': '',
		})
		self.assertTrue(form.is_valid(), form.errors)
		registration = form.save(commit=False)
		registration.student = self.student
		registration.save()
		self.assertEqual(registration.get_group_display(), 'Science')


class AbsentSubjectRulesTests(TestCase):
	"""NCTB/SSC reading: not sitting an assigned subject is a fail, not a free pass."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Absent School', classes='9')
		self.user = get_user_model().objects.create_superuser(username='absent-admin', password='password')
		self.client.force_login(self.user)
		self.physics = Subject.objects.create(code='PHYA', name='Physics', full_marks=100)
		self.math = Subject.objects.create(code='MTXA', name='Higher Math', full_marks=100)
		from .models import SubjectRequirement
		for subject in (self.physics, self.math):
			# Group-neutral assignments, so both subjects are columns even when a
			# student has no mark in one of them.
			SubjectRequirement.objects.create(
				institution=self.institution, admission_class='9', group='',
				subject=subject, requirement_type='MANDATORY',
			)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='9', session='2026', is_published=True,
		)
		self.partial = Student.objects.create(
			institution=self.institution, student_id='A001', name='Half Present',
			admission_class='9', section='A', roll_no=1, admission_year=2026,
		)
		self.never = Student.objects.create(
			institution=self.institution, student_id='A002', name='Totally Absent',
			admission_class='9', section='A', roll_no=2, admission_year=2026,
		)
		ExamMark.objects.create(
			exam=self.exam, student=self.partial, subject=self.physics, marks_obtained=95,
		)

	def test_unentered_subject_is_graded_f_and_makes_the_result_fail(self):
		result = compute_first_result(self.exam)
		self.assertEqual(result['student'], self.partial)
		self.assertEqual(result['status'], 'Fail')
		self.assertEqual(str(result['gpa']), '0.00')
		absent_row = [row for row in result['subject_results'] if row['absent']][0]
		self.assertEqual(absent_row['subject'], self.math)
		self.assertEqual(absent_row['grade'], 'F')
		self.assertEqual(absent_row['obtained'], 0)
		# the missed subject is counted as 0 / 100, so the percentage tells the truth
		self.assertEqual(result['total_obtained'], 95)
		self.assertEqual(result['total_full'], 200)

	def test_a_dash_is_printed_for_the_absent_subject_not_a_zero(self):
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		content = response.content.decode()
		self.assertIn('absent-mark', content)
		self.assertIn('no mark entered - counted as F', content)

	def test_student_with_no_marks_anywhere_is_no_marks_not_fail(self):
		"""Nobody sat this exam: an attendance problem, not a graded result."""
		result = [r for r in self._results() if r['student'] == self.never][0]
		self.assertEqual(result['status'], 'No Marks')
		self.assertIsNone(result['gpa'])
		self.assertIsNone(result['position'])

	def test_absent_rule_can_be_switched_off(self):
		from django.test import override_settings
		with override_settings(EXAM_ABSENT_SUBJECT_FAILS=False):
			result = compute_first_result(self.exam)
		self.assertEqual(result['status'], 'Pass')
		self.assertEqual(result['total_full'], 100)
		absent_row = [row for row in result['subject_results'] if row['absent']][0]
		self.assertEqual(absent_row['grade'], '-')
		self.assertIsNone(absent_row['obtained'])

	def _results(self):
		from .result_utils import build_exam_results
		_, results = build_exam_results(self.exam)
		return results


class GroupOnlyFromClass9Tests(TestCase):
	"""A group exists from class 9 up. Primary and junior secondary (6-8)
	follow one common syllabus, so the field must be hidden, optional and
	never stored for them."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Group Rule Campus', classes='6,7,8,9,10,11,12')
		self.user = get_user_model().objects.create_superuser(username='group-admin', password='password')
		self.client.force_login(self.user)

	def _post(self, **overrides):
		data = {
			'institution': self.institution.pk,
			'name': 'Group Rule Student',
			'admission_class': '6',
			'section': 'A',
			'admission_year': 2026,
			'roll_no': 1,
			'gender': 'M',
			'religion': 'Islam',
			'status': 'ACTIVE',
			'group': 'SCI',
		}
		data.update(overrides)
		return self.client.post(reverse('add_student'), data)

	def test_group_choices_start_at_class_9(self):
		self.assertFalse(Student.class_supports_group('6'))
		self.assertFalse(Student.class_supports_group('08'))
		self.assertTrue(Student.class_supports_group('9'))
		self.assertTrue(Student.class_supports_group('12'))
		codes = [code for code, _ in Student.group_choices_for_class('9')]
		self.assertEqual(codes, ['SCI', 'BUS', 'HUM'])
		self.assertEqual(Student.group_choices_for_class('6'), [])

	def test_class_6_admission_ignores_a_submitted_group(self):
		response = self._post(admission_class='6', group='SCI')
		self.assertEqual(response.status_code, 302)
		student = Student.objects.get(name='Group Rule Student')
		self.assertEqual(student.group, '')

	def test_class_9_admission_requires_a_group(self):
		response = self._post(name='No Group Nine', admission_class='9', group='')
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Group is required')
		self.assertFalse(Student.objects.filter(name='No Group Nine').exists())

	def test_class_9_admission_accepts_business_and_humanities(self):
		for code in ('BUS', 'HUM'):
			self._post(name=f'Student {code}', admission_class='9', group=code)
			self.assertEqual(Student.objects.get(name=f'Student {code}').group, code)

	def test_diploma_group_is_not_offered_at_class_9(self):
		"""A stale dropdown can still post a code the class does not offer;
		the officer gets a usable message, not a bare 'invalid choice'."""
		response = self._post(name='Diploma Nine', admission_class='9', group='DCS')
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Select Science, Business Studies or Humanities.')
		self.assertFalse(Student.objects.filter(name='Diploma Nine').exists())

	def test_model_save_clears_a_stale_group(self):
		student = Student.objects.create(
			institution=self.institution, student_id='G001', name='Stale Group',
			admission_class='7', section='A', admission_year=2026, group='SCI',
		)
		self.assertEqual(student.group, '')
		student.admission_class = '9'
		student.group = 'HUM'
		student.save()
		self.assertEqual(student.group, 'HUM')

	def test_add_student_page_hides_the_group_field_below_class_9(self):
		response = self.client.get(reverse('add_student'))
		self.assertContains(response, 'id="group-field-wrapper"')
		self.assertContains(response, 'const GROUPED_CLASSES = ["9", "10", "11", "12"];')

	def test_admission_dropdown_offers_no_group_below_class_9(self):
		response = self.client.get(reverse('admission_dropdown_options'), {
			'institution': self.institution.pk, 'admission_class': '6',
		})
		payload = response.json()
		self.assertEqual(payload['groups'], [])
		self.assertFalse(payload['supports_group'])

	def test_admission_dropdown_offers_three_groups_from_class_9(self):
		for cls in ('9', '11'):
			payload = self.client.get(reverse('admission_dropdown_options'), {
				'institution': self.institution.pk, 'admission_class': cls,
			}).json()
			self.assertTrue(payload['supports_group'])
			self.assertEqual([g['value'] for g in payload['groups']], ['SCI', 'BUS', 'HUM'])

	def test_bulk_update_into_class_6_clears_the_group(self):
		student = Student.objects.create(
			institution=self.institution, student_id='G002', name='Moving Down',
			admission_class='9', section='A', admission_year=2026, group='SCI',
		)
		response = self.client.post(reverse('bulk_update_students'), {
			'student_ids': [student.pk], 'new_class': '6', 'new_group': 'HUM',
		})
		self.assertEqual(response.status_code, 302)
		student.refresh_from_db()
		self.assertEqual(student.admission_class, '6')
		self.assertEqual(student.group, '')

	def test_bulk_update_keeps_the_group_inside_class_9(self):
		student = Student.objects.create(
			institution=self.institution, student_id='G003', name='Staying Nine',
			admission_class='9', section='A', admission_year=2026, group='SCI',
		)
		self.client.post(reverse('bulk_update_students'), {
			'student_ids': [student.pk], 'new_group': 'BUS',
		})
		student.refresh_from_db()
		self.assertEqual(student.group, 'BUS')

	@skipUnless(Workbook, 'openpyxl not installed')
	def test_excel_import_ignores_the_group_column_below_class_9(self):
		book = Workbook()
		sheet = book.active
		sheet.append([
			'Institution', 'Name', 'Class', 'Section', 'Admission Year', 'Roll No',
			'Gender', 'Religion', 'Father Name', 'Contact No', 'Guardian Contact No', 'Group',
		])
		sheet.append([
			self.institution.name, 'Import Six', '6', 'A', 2026, 11,
			'Male', 'Islam', 'Father', '', '', 'Science',
		])
		sheet.append([
			self.institution.name, 'Import Nine', '9', 'A', 2026, 12,
			'Female', 'Islam', 'Father', '', '', 'Humanities',
		])
		buffer = BytesIO()
		book.save(buffer)
		buffer.seek(0)

		response = self.client.post(reverse('import_students'), {
			'excel_file': SimpleUploadedFile('students.xlsx', buffer.getvalue()),
		})
		self.assertEqual(response.status_code, 302)  # redirects back to the list on success
		self.assertEqual(Student.objects.get(name='Import Six').group, '')
		self.assertEqual(Student.objects.get(name='Import Nine').group, 'HUM')

	def test_clean_student_groups_command_only_touches_classes_below_9(self):
		from io import StringIO
		from django.core.management import call_command

		junior = Student.objects.create(
			institution=self.institution, student_id='G004', name='Junior Wrong',
			admission_class='6', section='A', admission_year=2026,
		)
		# bypasses save() on purpose: this is the state the live data is in
		Student.objects.filter(pk=junior.pk).update(group='SCI')
		senior = Student.objects.create(
			institution=self.institution, student_id='G005', name='Senior Right',
			admission_class='9', section='A', admission_year=2026, group='SCI',
		)

		out = StringIO()
		call_command('clean_student_groups', stdout=out)
		junior.refresh_from_db()
		self.assertEqual(junior.group, 'SCI', 'dry run must not change anything')
		self.assertIn('Dry run', out.getvalue())

		out = StringIO()
		call_command('clean_student_groups', '--apply', stdout=out)
		junior.refresh_from_db()
		senior.refresh_from_db()
		self.assertEqual(junior.group, '')
		self.assertEqual(senior.group, 'SCI')
		self.assertIn('Cleared the group on 1 student(s)', out.getvalue())

	def test_data_migration_clears_groups_below_class_9(self):
		# The migration module name starts with a digit, so it cannot be
		# imported with a normal import statement — load it by path instead.
		import importlib.util
		import pathlib

		path = pathlib.Path(__file__).parent / 'migrations' / '0033_clear_groups_below_class_9.py'
		spec = importlib.util.spec_from_file_location('clear_groups_migration', path)
		migration = importlib.util.module_from_spec(spec)
		spec.loader.exec_module(migration)

		wrong = Student.objects.create(
			institution=self.institution, student_id='G006', name='Migration Target',
			admission_class='8', section='A', admission_year=2026,
		)
		Student.objects.filter(pk=wrong.pk).update(group='HUM')

		class AppsShim:
			"""Stands in for the migration's historical app registry."""
			@staticmethod
			def get_model(app_label, model_name):
				return Student

		migration.clear_groups_below_class_9(AppsShim, None)
		wrong.refresh_from_db()
		self.assertEqual(wrong.group, '')


class GroupRuleTemplateSmokeTests(TestCase):
    """The group-rule templates were edited by hand; make sure they render."""

    def setUp(self):
        self.institution = Institution.objects.create(name='Smoke Campus', classes='6,9')
        self.user = get_user_model().objects.create_superuser(username='smoke', password='pw')
        self.client.force_login(self.user)
        self.six = Student.objects.create(
            institution=self.institution, student_id='S001', name='Six Student',
            admission_class='6', section='A', admission_year=2026,
        )

    def test_pages_render(self):
        for name, kwargs in [
            ('add_student', {}),
            ('admission', {}),
            ('public_admission_apply', {}),
            ('student_list', {}),
        ]:
            response = self.client.get(reverse(name, kwargs=kwargs))
            self.assertEqual(response.status_code, 200, name)

        response = self.client.post(reverse('bulk_update_select'), {
            'student_ids': [self.six.pk], 'institution': self.institution.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'GROUPED_CLASSES')

        response = self.client.get(reverse('edit_student', args=[self.six.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="group-field-wrapper"')

    def test_add_student_edit_page_marks_the_group_field_hidden_for_class_6(self):
        response = self.client.get(reverse('edit_student', args=[self.six.pk]))
        self.assertContains(response, 'const GROUPED_CLASSES = ["9", "10", "11", "12"];')
        self.assertContains(response, 'This class has no group')

    def test_admission_application_form_hides_the_group_until_a_class_is_picked(self):
        response = self.client.get(reverse('admission'))
        self.assertContains(response, 'function setGroupFieldVisible(visible, note)')
        self.assertContains(response, 'setGroupFieldVisible(false, \'\');')
        self.assertContains(response, 'data.supports_group')

    def test_public_admission_form_marks_the_group_column(self):
        response = self.client.get(reverse('public_admission_apply'))
        self.assertContains(response, 'id="group-field-wrapper"')
        self.assertContains(response, 'id="group-field-note"')


class InstitutionAccessSetupTests(TestCase):
	"""The login screen must work on a fresh install where nobody has been
	granted an InstitutionAccess row yet."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Kallan Trust Campus', classes='6,7,8,9')
		self.other = Institution.objects.create(name='Second Campus', classes='6,7')

	def test_login_page_lists_institutions_and_departments_without_any_access_rows(self):
		self.assertEqual(InstitutionAccess.objects.count(), 0)
		response = self.client.get(reverse('login'))
		self.assertEqual(response.status_code, 200)
		content = response.content.decode()
		self.assertContains(response, 'Kallan Trust Campus')
		self.assertContains(response, 'Second Campus')
		for department in ('Office', 'Exam', 'Accounts'):
			self.assertContains(response, f'data-department="{department}"')
		self.assertNotContains(response, 'No institution access has been assigned')
		# the form must post a real institution, not an empty one
		self.assertIn(f'name="institution_id" id="institution_id" value="{self.institution.pk}"', content)

	def test_admin_can_log_in_with_a_fresh_install(self):
		get_user_model().objects.create_superuser(username='boss', password='secret123')
		response = self.client.post(reverse('login'), {
			'username': 'boss', 'password': 'secret123',
			'institution_id': str(self.institution.pk), 'department': 'Office',
		})
		self.assertEqual(response.status_code, 302)
		self.assertEqual(self.client.session['selected_institution_id'], str(self.institution.pk))
		self.assertEqual(self.client.session['selected_department'], 'Office')

	def test_plain_user_without_access_is_still_refused(self):
		get_user_model().objects.create_user(username='stranger', password='secret123')
		response = self.client.post(reverse('login'), {
			'username': 'stranger', 'password': 'secret123',
			'institution_id': str(self.institution.pk), 'department': 'Office',
		})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Invalid username, password, or institution access.')

	def test_grant_command_lets_a_plain_user_in(self):
		from io import StringIO
		from django.core.management import call_command

		get_user_model().objects.create_user(username='clerk', password='secret123')

		# before the grant
		response = self.client.post(reverse('login'), {
			'username': 'clerk', 'password': 'secret123',
			'institution_id': str(self.institution.pk), 'department': 'Office',
		})
		self.assertContains(response, 'Invalid username, password, or institution access.')

		out = StringIO()
		call_command('grant_institution_access', 'clerk',
		             '--institution', str(self.institution.pk), '--department', 'Office',
		             stdout=out)
		self.assertIn('can now log in', out.getvalue())
		self.assertTrue(InstitutionAccess.objects.filter(
			user__username='clerk', institution=self.institution, department='Office', is_active=True,
		).exists())

		response = self.client.post(reverse('login'), {
			'username': 'clerk', 'password': 'secret123',
			'institution_id': str(self.institution.pk), 'department': 'Office',
		})
		self.assertEqual(response.status_code, 302)
		self.assertEqual(self.client.session['selected_institution_id'], str(self.institution.pk))
		# the department group came with the grant, so the pages are not 403
		user = get_user_model().objects.get(username='clerk')
		self.assertIn('Office', [g.name for g in user.groups.all()])

	def test_grant_command_accepts_an_institution_name_and_is_idempotent(self):
		from io import StringIO
		from django.core.management import call_command

		user = get_user_model().objects.create_user(username='clerk2', password='secret123')
		for _ in range(2):
			call_command('grant_institution_access', 'clerk2',
			             '--institution', 'Kallan Trust Campus', stdout=StringIO())
		self.assertEqual(InstitutionAccess.objects.filter(user=user, institution=self.institution).count(), 3)

	def test_revoke_command_takes_the_access_away(self):
		from io import StringIO
		from django.core.management import call_command

		user = get_user_model().objects.create_user(username='clerk3', password='secret123')
		call_command('grant_institution_access', 'clerk3', '--institution', str(self.institution.pk),
		             '--department', 'Exam', stdout=StringIO())
		out = StringIO()
		call_command('grant_institution_access', 'clerk3', '--institution', str(self.institution.pk),
		             '--department', 'Exam', '--revoke', stdout=out)
		self.assertIn('Revoked 1 access row(s)', out.getvalue())
		self.assertFalse(InstitutionAccess.objects.filter(user=user).exists())


class ArchiveIntegrityTests(TestCase):
	"""An archived student is out of every active surface — reports, exports,
	promotion, bulk edits — and the people who archive must be able to read the
	archive back."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Archive Campus', classes='6,7,8,9')
		self.admin = get_user_model().objects.create_superuser(username='arch-admin', password='pw')
		self.client.force_login(self.admin)
		self.archived = Student.objects.create(
			institution=self.institution, student_id='AR001', name='Left The School',
			admission_class='6', section='A', admission_year=2026, roll_no=1,
		)
		self.active = Student.objects.create(
			institution=self.institution, student_id='AR002', name='Still Here',
			admission_class='6', section='A', admission_year=2026, roll_no=2,
		)
		self.client.post(reverse('delete_student', args=[self.archived.pk]))
		self.archived.refresh_from_db()
		self.assertTrue(self.archived.is_archived)

	def _office_user(self):
		from .permissions import ensure_default_groups
		ensure_default_groups()
		user = get_user_model().objects.create_user(username='office_arch', password='pw')
		InstitutionAccess.objects.create(user=user, institution=self.institution, department='Office')
		self.client.logout()
		self.client.post(reverse('login'), {
			'username': 'office_arch', 'password': 'pw',
			'institution_id': str(self.institution.pk), 'department': 'Office',
		})
		return user

	def test_promotion_leaves_archived_students_behind(self):
		response = self.client.post(reverse('student_promotion'), {
			'from_class': '6', 'from_section': 'A',
			'to_class': '7', 'to_section': 'A', 'session': '2026-2027',
		})
		self.assertEqual(response.status_code, 302)
		self.archived.refresh_from_db()
		self.active.refresh_from_db()
		self.assertEqual(self.active.admission_class, '7')
		self.assertEqual(self.archived.admission_class, '6', 'an archived student must not be promoted')

	def test_excel_export_matches_the_list_and_skips_archived(self):
		response = self.client.get(
			reverse('download_student_list') + f'?institution={self.institution.pk}&all=1'
		)
		self.assertEqual(response.status_code, 200)
		from openpyxl import load_workbook
		sheet = load_workbook(BytesIO(response.content)).active
		names = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
		self.assertIn('Still Here', names)
		self.assertNotIn('Left The School', names)

	def test_class_section_summary_counts_only_active_students(self):
		response = self.client.get(
			reverse('class_section_summary') + f'?institution={self.institution.pk}'
		)
		body = response.content.decode()
		self.assertNotIn('Left The School', body)
		# class 6 / section A now holds exactly one active student
		self.assertContains(response, '<td>6</td>', html=False)

	def test_bulk_update_skips_archived_students(self):
		self.client.post(reverse('bulk_update_students'), {
			'student_ids': [self.archived.pk, self.active.pk], 'new_section': 'Z',
		})
		self.archived.refresh_from_db()
		self.active.refresh_from_db()
		self.assertEqual(self.active.section, 'Z')
		self.assertEqual(self.archived.section, 'A', 'bulk update must not touch the archive')

	def test_auto_registration_skips_archived_students(self):
		from .curriculum_apply import apply_curriculum
		apply_curriculum(self.institution, '6')
		response = self.client.post(reverse('auto_register_students'), {
			'student_ids': [self.archived.pk],
		})
		self.assertEqual(response.status_code, 302)
		self.assertEqual(StudentSubjectChoice.objects.filter(student=self.archived).count(), 0)

	def test_editing_an_archived_student_sends_you_to_the_archive(self):
		response = self.client.get(reverse('edit_student', args=[self.archived.pk]))
		self.assertEqual(response.status_code, 302)
		self.assertIn(reverse('archived_students'), response.url)

	def test_office_department_can_read_the_archive_it_created(self):
		user = self._office_user()
		self.assertIn('students.view_student', user.get_all_permissions())
		response = self.client.get(reverse('archived_students'))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Left The School')

	def test_office_department_can_restore_what_it_archived(self):
		self._office_user()
		response = self.client.post(reverse('restore_student', args=[self.archived.pk]))
		self.assertEqual(response.status_code, 302)
		self.archived.refresh_from_db()
		self.assertFalse(self.archived.is_archived)
		self.assertEqual(self.archived.status, 'ACTIVE')

	def test_restore_needs_the_same_permission_as_archive(self):
		"""A user who may change students but not archive them must not be able
		to undo an archive. No InstitutionAccess row here on purpose — that is
		what keeps the department groups (and their delete_student) away."""
		from django.contrib.auth.models import Group
		group = Group.objects.create(name='Editors Only')
		ct = ContentType.objects.get_for_model(Student)
		group.permissions.add(Permission.objects.get(content_type=ct, codename='change_student'))
		user = get_user_model().objects.create_user(username='editor', password='pw')
		user.groups.add(group)
		self.client.logout()
		self.client.force_login(user)
		self.assertNotIn('students.delete_student', user.get_all_permissions())

		for view_name, payload in [
			('bulk_restore_students', {'student_ids': [self.archived.pk]}),
			('bulk_purge_archived_students', {'student_ids': [self.archived.pk]}),
		]:
			response = self.client.post(reverse(view_name), payload)
			self.assertEqual(response.status_code, 403, view_name)
		response = self.client.post(reverse('restore_student', args=[self.archived.pk]))
		self.assertEqual(response.status_code, 403)
		response = self.client.post(reverse('purge_archived_student', args=[self.archived.pk]))
		self.assertEqual(response.status_code, 403)
		self.archived.refresh_from_db()
		self.assertTrue(self.archived.is_archived)
		self.assertTrue(Student.objects.filter(pk=self.archived.pk).exists())

	@skipUnless(Workbook, 'openpyxl not installed')
	def test_ssc_import_will_not_attach_to_an_archived_student(self):
		book = Workbook()
		sheet = book.active
		sheet.append(['Student ID', 'Registration No', 'Roll No', 'Session', 'Group', 'Subjects', 'Board'])
		sheet.append(['AR001', 'REG-900', 1, '2025-2026', 'Science', '', 'Dhaka'])
		buffer = BytesIO()
		book.save(buffer)
		buffer.seek(0)
		self.client.post(reverse('import_ssc_registrations'), {
			'excel_file': SimpleUploadedFile('regs.xlsx', buffer.getvalue()),
		})
		self.assertFalse(SSCRegistration.objects.filter(student=self.archived).exists())


class StudentImportLabelTests(TestCase):
	"""The filled import template is written by hand: 'Business' not
	'Business Studies'. The import must survive that, and re-running the same
	file must not duplicate anyone."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Import Campus', classes='6,9,10')
		self.admin = get_user_model().objects.create_superuser(username='import-admin', password='pw')
		self.client.force_login(self.admin)

	def _upload(self, rows):
		book = Workbook()
		sheet = book.active
		sheet.append([
			'Institution', 'Name', 'Class', 'Section', 'Admission Year', 'Roll No',
			'Gender', 'Religion', 'Father Name', 'Contact No', 'Guardian Contact No', 'Group',
		])
		for row in rows:
			sheet.append(row)
		buffer = BytesIO()
		book.save(buffer)
		buffer.seek(0)
		return self.client.post(reverse('import_students'), {
			'excel_file': SimpleUploadedFile('students.xlsx', buffer.getvalue()),
		}, follow=True)

	def _row(self, name, roll, group, year=2026, cls='9'):
		return [self.institution.name, name, cls, 'A', year, roll,
		        'Male', 'Islam', 'Father', '', '', group]

	def test_hand_written_group_labels_are_mapped(self):
		response = self._upload([
			self._row('Science Kid', 1, 'Science'),
			self._row('Business Kid', 2, 'Business'),
			self._row('Humanities Kid', 3, 'Humanities'),
			self._row('Code Kid', 4, 'BUS'),
			self._row('No Group Kid', 5, ''),
		])
		self.assertEqual(response.status_code, 200)
		self.assertEqual(Student.objects.get(name='Science Kid').group, 'SCI')
		self.assertEqual(Student.objects.get(name='Business Kid').group, 'BUS')
		self.assertEqual(Student.objects.get(name='Humanities Kid').group, 'HUM')
		self.assertEqual(Student.objects.get(name='Code Kid').group, 'BUS')
		self.assertEqual(Student.objects.get(name='No Group Kid').group, '')

	def test_group_is_left_blank_below_class_9_even_when_the_sheet_says_science(self):
		self._upload([self._row('Junior Kid', 1, 'Science', cls='6')])
		self.assertEqual(Student.objects.get(name='Junior Kid').group, '')

	def test_re_importing_the_same_file_does_not_duplicate(self):
		rows = [self._row('Again Kid', 7, 'Business'), self._row('Again Two', 8, 'Humanities')]
		self._upload(rows)
		self._upload(rows)
		self.assertEqual(Student.objects.filter(name='Again Kid').count(), 1)
		self.assertEqual(Student.objects.filter(name='Again Two').count(), 1)
		self.assertEqual(Student.objects.filter(institution=self.institution).count(), 2)

	def test_re_import_fills_a_missing_group_on_the_existing_student(self):
		# first file had no group column value; the student was created blank
		self._upload([self._row('Blank Group Kid', 9, '')])
		student = Student.objects.get(name='Blank Group Kid')
		self.assertEqual(student.group, '')
		# second file carries the group; the re-run must backfill it
		response = self._upload([self._row('Blank Group Kid', 9, 'Business')])
		messages = [str(m) for m in response.context['messages']]
		student.refresh_from_db()
		self.assertEqual(student.group, 'BUS')
		self.assertTrue(any('1 existing student(s) got their group filled in' in m for m in messages))


class StudentIdHoleTests(TestCase):
	"""A partial import leaves holes; the next import must step over the taken
	ids instead of dying on an IntegrityError."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Hole Campus', classes='9')

	def test_new_student_skips_taken_suffixes(self):
		# holes on purpose: suffixes 1 and 3 exist, 2 is free, then a gap
		Student.objects.create(institution=self.institution, student_id='202509001',
		                       name='One', admission_class='9', section='A', admission_year=2025)
		Student.objects.create(institution=self.institution, student_id='202509003',
		                       name='Three', admission_class='9', section='A', admission_year=2025)
		new_one = Student.objects.create(institution=self.institution, name='New A',
		                                 admission_class='9', section='A', admission_year=2025)
		new_two = Student.objects.create(institution=self.institution, name='New B',
		                                 admission_class='9', section='A', admission_year=2025)
		ids = {new_one.student_id, new_two.student_id}
		self.assertNotIn('202509001', ids)
		self.assertNotIn('202509003', ids)
		self.assertEqual(len(ids), 2)


class StudentListCountAndLookupTests(TestCase):
	"""The list must show how many students are on screen, and an ID lookup
	must open that student's profile — otherwise a 132/133 gap is invisible."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Count Campus', classes='6,9')
		self.user = get_user_model().objects.create_superuser(username='count-admin', password='pw')
		self.client.force_login(self.user)
		self.six_a = Student.objects.create(
			institution=self.institution, student_id='CNT001', name='Six A One',
			admission_class='6', section='A', roll_no=1, admission_year=2026,
		)
		self.six_b = Student.objects.create(
			institution=self.institution, student_id='CNT002', name='Six B One',
			admission_class='6', section='B', roll_no=1, admission_year=2026,
		)
		self.nine = Student.objects.create(
			institution=self.institution, student_id='CNT009', name='Nine Science',
			admission_class='9', section='A', group='SCI', roll_no=1, admission_year=2026,
		)

	def test_list_shows_total_and_class_section_counts(self):
		response = self.client.get(reverse('student_list'), {'institution': self.institution.pk})
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['student_count'], 3)
		self.assertContains(response, 'id="student-count"')
		self.assertContains(response, '>3</strong>')
		body = response.content.decode()
		self.assertIn('Class 6 A', body)
		self.assertIn('Class 6 B', body)
		self.assertIn('Class 9 A', body)
		self.assertContains(response, 'Science')

	def test_exact_student_id_opens_the_profile(self):
		response = self.client.get(reverse('student_list'), {'q': 'CNT009'})
		self.assertRedirects(response, reverse('student_detail', args=[self.nine.pk]))

		response = self.client.get(reverse('student_by_id', args=['CNT001']))
		self.assertRedirects(response, reverse('student_detail', args=[self.six_a.pk]))

	def test_unknown_id_stays_on_the_list_with_a_message(self):
		response = self.client.get(reverse('student_by_id', args=['MISSING99']), follow=True)
		self.assertContains(response, 'MISSING99')
		self.assertContains(response, 'No student matching')

	def test_name_search_filters_the_list_without_redirecting(self):
		response = self.client.get(reverse('student_list'), {'q': 'Six'})
		self.assertEqual(response.status_code, 200)
		self.assertEqual(response.context['student_count'], 2)
		self.assertContains(response, 'Six A One')
		self.assertContains(response, 'Six B One')
		self.assertNotContains(response, 'Nine Science')

	def test_archived_student_is_still_reachable_by_id(self):
		self.six_a.is_archived = True
		self.six_a.save(update_fields=['is_archived'])
		response = self.client.get(reverse('student_list'), {'q': 'CNT001'})
		self.assertRedirects(response, reverse('student_detail', args=[self.six_a.pk]))
		profile = self.client.get(reverse('student_detail', args=[self.six_a.pk]))
		self.assertContains(profile, 'archived')


class ReligionPaperTests(TestCase):
	"""The two religion papers print as ONE merged Religion column (REL).
	Each student is graded on their own paper (Hindu students sit Hindu
	Religion & Moral Education, everyone else Islam & Moral Education); a
	student whose own paper is not assigned gets a dash that is never failed
	or counted. Papers for religions the school has no students of never
	appear."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Religion School', classes='6,9')
		self.user = get_user_model().objects.create_superuser(username='religion-admin', password='password')
		self.client.force_login(self.user)
		from .models import SubjectRequirement
		self.bangla = Subject.objects.create(code='RBN', name='Bangla', full_marks=100)
		self.islam = Subject.objects.create(
			code='RISL', name='Islam & Moral Education', full_marks=100, category='RELIGION')
		self.hindu = Subject.objects.create(
			code='RHIN', name='Hindu Religion & Moral Education', full_marks=100, category='RELIGION')
		self.christian = Subject.objects.create(
			code='RCHR', name='Christian Religion & Moral Education', full_marks=100, category='RELIGION')
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='6', subject=self.bangla,
			requirement_type='MANDATORY',
		)
		for subject, religion in (
			(self.islam, 'Islam'), (self.hindu, 'Hindu'), (self.christian, 'Christian'),
		):
			SubjectRequirement.objects.create(
				institution=self.institution, admission_class='6', subject=subject,
				requirement_type='CONDITIONAL', condition_religion=religion,
			)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='6', session='2026', is_published=True,
		)
		self.muslim = Student.objects.create(
			institution=self.institution, student_id='R001', name='Muslim Kid',
			admission_class='6', section='A', roll_no=1, admission_year=2026, religion='Islam',
		)
		self.hindu_kid = Student.objects.create(
			institution=self.institution, student_id='R002', name='Hindu Kid',
			admission_class='6', section='A', roll_no=2, admission_year=2026, religion='Hindu',
		)

	def _marks(self, student, subject, value):
		ExamMark.objects.create(
			exam=self.exam, student=student, subject=subject, marks_obtained=value,
		)

	def _built(self):
		from .result_utils import build_exam_results
		return build_exam_results(self.exam)

	def _result_for(self, student):
		_, results = self._built()
		return next(r for r in results if r['student'].pk == student.pk)

	def _religion_row(self, result):
		return next(
			row for row in result['subject_results'] if row.get('religion_column')
		)

	def test_the_two_papers_print_as_one_religion_column(self):
		from .result_utils import ReligionColumn
		columns, _results = self._built()
		religion_columns = [column for column in columns if isinstance(column, ReligionColumn)]
		self.assertEqual(len(religion_columns), 1)
		religion_column = religion_columns[0]
		self.assertEqual(religion_column.code, 'REL')
		# Neither underlying paper is its own column any more.
		self.assertNotIn(self.islam, columns)
		self.assertNotIn(self.hindu, columns)
		# A paper nobody sits (Christian) is not part of the merged column.
		self.assertNotIn(self.christian, columns)
		self.assertNotIn(self.christian, religion_column.subjects)
		self.assertIn(self.islam, religion_column.subjects)
		self.assertIn(self.hindu, religion_column.subjects)
		# Bangla stays a normal column.
		self.assertIn(self.bangla, columns)

	def test_each_student_is_graded_on_their_own_religion_paper(self):
		self._marks(self.muslim, self.bangla, 80)
		self._marks(self.muslim, self.islam, 75)
		self._marks(self.hindu_kid, self.bangla, 82)
		self._marks(self.hindu_kid, self.hindu, 71)

		muslim_result = self._result_for(self.muslim)
		muslim_rel = self._religion_row(muslim_result)
		self.assertEqual(muslim_rel['paper'], self.islam)
		self.assertEqual(muslim_rel['obtained'], 75)
		self.assertEqual(muslim_result['status'], 'Pass')
		self.assertEqual(muslim_result['total_full'], 200)  # Bangla + Islam only
		self.assertEqual(muslim_result['position'], 1)

		hindu_result = self._result_for(self.hindu_kid)
		hindu_rel = self._religion_row(hindu_result)
		self.assertEqual(hindu_rel['paper'], self.hindu)
		self.assertEqual(hindu_rel['obtained'], 71)
		self.assertEqual(hindu_result['status'], 'Pass')
		self.assertEqual(hindu_result['total_full'], 200)  # Bangla + Hindu only
		self.assertEqual(hindu_result['position'], 2)

	def test_religion_column_is_one_row_in_detail_and_card(self):
		self._marks(self.muslim, self.bangla, 80)
		self._marks(self.muslim, self.islam, 75)
		self._marks(self.hindu_kid, self.bangla, 82)
		self._marks(self.hindu_kid, self.hindu, 71)
		for url_name in ('student_result_detail', 'result_card'):
			response = self.client.get(
				reverse(url_name, args=[self.exam.pk, self.hindu_kid.pk])
			)
			self.assertEqual(response.status_code, 200, url_name)
			content = response.content.decode()
			# Exactly one Religion subject row (the merged column).
			self.assertEqual(content.count('<td>Religion'), 1, url_name)
			# The other religion paper's name never appears as a subject row.
			self.assertNotIn('<td>Islam &amp; Moral Education</td>', content)
			self.assertNotIn('<td>Hindu Religion &amp; Moral Education</td>', content)

	def test_a_hindu_student_without_a_hindu_paper_gets_a_never_counted_dash(self):
		# Only the Islam paper is assigned: the Hindu student has no own paper,
		# so the Religion cell is a dash — never a fail, never counted, and the
		# Islam paper is never forced on them.
		from .models import SubjectRequirement
		SubjectRequirement.objects.filter(subject__in=[self.hindu, self.christian]).delete()
		self._marks(self.muslim, self.bangla, 80)
		self._marks(self.muslim, self.islam, 75)
		self._marks(self.hindu_kid, self.bangla, 82)
		# A stray Islam mark for the Hindu student must be ignored.
		self._marks(self.hindu_kid, self.islam, 70)

		hindu_result = self._result_for(self.hindu_kid)
		rel_row = self._religion_row(hindu_result)
		self.assertTrue(rel_row['religion_unassigned'])
		self.assertTrue(rel_row['absent'])
		self.assertIsNone(rel_row['obtained'])
		self.assertIsNone(rel_row['point'])
		# Only Bangla counts: the dash never fails or inflates the total.
		self.assertEqual(hindu_result['total_full'], 100)
		self.assertEqual(hindu_result['total_obtained'], 82)
		self.assertEqual(hindu_result['status'], 'Pass')
		self.assertEqual(hindu_result['absent_subject_count'], 0)

		# The Muslim student still sits Islam through the merged column.
		muslim_result = self._result_for(self.muslim)
		muslim_rel = self._religion_row(muslim_result)
		self.assertEqual(muslim_rel['paper'], self.islam)
		self.assertEqual(muslim_result['status'], 'Pass')

	def test_no_religion_papers_assigned_means_no_religion_column(self):
		from .models import SubjectRequirement
		SubjectRequirement.objects.filter(subject__category='RELIGION').delete()
		columns, results = self._built()
		self.assertFalse(any(getattr(c, 'code', '') == 'REL' for c in columns))
		result = self._result_for(self.muslim)
		self.assertFalse(any(r.get('religion_column') for r in result['subject_results']))

	def test_blank_religion_sits_the_islam_paper(self):
		blank = Student.objects.create(
			institution=self.institution, student_id='R003', name='No Religion Set',
			admission_class='6', section='A', roll_no=3, admission_year=2026, religion='',
		)
		self._marks(blank, self.bangla, 60)
		self._marks(blank, self.islam, 55)
		result = self._result_for(blank)
		rel_row = self._religion_row(result)
		self.assertEqual(rel_row['paper'], self.islam)
		self.assertEqual(result['status'], 'Pass')
		self.assertEqual(result['total_full'], 200)

	def test_result_sheet_shows_rel_code_and_legend(self):
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		self.assertEqual(response.status_code, 200)
		content = response.content.decode()
		self.assertIn('Subject codes', content)
		self.assertIn('<span class="code">REL</span> = Religion &amp; Moral Education', content)
		self.assertIn('student-col', content)
		# The header uses the code, not the long paper names.
		self.assertNotIn('<th>Islam &amp; Moral Education', content)
		self.assertNotIn('<th>Hindu Religion &amp; Moral Education', content)

	def test_marks_entry_only_offers_papers_students_sit(self):
		from .result_utils import get_exam_subjects_for_students, get_exam_students
		students = list(get_exam_students(self.exam))
		subjects, is_filtered = get_exam_subjects_for_students(self.exam, students)
		self.assertTrue(is_filtered)
		self.assertIn(self.islam, subjects)
		self.assertIn(self.hindu, subjects)
		self.assertNotIn(self.christian, subjects)

	def test_enter_marks_skips_students_who_do_not_sit_the_paper(self):
		# Entering Islam marks: the Hindu student's (disabled) boxes post
		# nothing and must never be saved, even when values sneak in.
		self.client.post(
			reverse('enter_marks', args=[self.exam.pk, self.islam.pk]),
			{
				f'marks_{self.muslim.pk}': '70',
				f'marks_{self.hindu_kid.pk}': '88',
			},
		)
		self.assertTrue(ExamMark.objects.filter(
			exam=self.exam, student=self.muslim, subject=self.islam).exists())
		self.assertFalse(ExamMark.objects.filter(
			exam=self.exam, student=self.hindu_kid, subject=self.islam).exists())


class AdmissionSubjectScopeTests(TestCase):
	"""Marks and results use the subjects selected for each admitted student,
	while retaining mandatory subjects and a union of subjects for the class
	columns."""

	def setUp(self):
		from .models import SubjectRequirement
		self.institution = Institution.objects.create(name='Admission Scope School', classes='9')
		self.user = get_user_model().objects.create_superuser(
			username='admission-scope-admin', password='password'
		)
		self.client.force_login(self.user)
		self.bangla = Subject.objects.create(code='ASB', name='Bangla', full_marks=100)
		self.ict = Subject.objects.create(code='ASI', name='ICT', full_marks=100)
		self.agriculture = Subject.objects.create(code='ASA', name='Agriculture', full_marks=100)
		self.bangla_requirement = SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', subject=self.bangla,
			requirement_type='MANDATORY',
		)
		self.ict_requirement = SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', subject=self.ict,
			requirement_type='OPTIONAL', optional_set_key='elective',
		)
		self.agriculture_requirement = SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', subject=self.agriculture,
			requirement_type='OPTIONAL', optional_set_key='elective',
		)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='9', session='2026', is_published=True,
		)
		self.ict_student = Student.objects.create(
			institution=self.institution, student_id='AS001', name='ICT Student',
			admission_class='9', section='A', roll_no=1, admission_year=2026,
		)
		self.agriculture_student = Student.objects.create(
			institution=self.institution, student_id='AS002', name='Agriculture Student',
			admission_class='9', section='A', roll_no=2, admission_year=2026,
		)
		# Mandatory subjects are intentionally not stored as choices here: the
		# admission model derives them from SubjectRequirement. Only the selected
		# optional subject is student-specific.
		for student, optional_requirement in (
			(self.ict_student, self.ict_requirement),
			(self.agriculture_student, self.agriculture_requirement),
		):
			StudentSubjectChoice.objects.create(student=student, requirement=optional_requirement)

	def test_unselected_optional_is_not_offered_on_any_exam_screen(self):
		from .models import SubjectRequirement
		higher_math = Subject.objects.create(code='ASH', name='Higher Math', full_marks=100)
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', subject=higher_math,
			requirement_type='OPTIONAL', optional_set_key='elective',
		)
		from .result_utils import get_exam_subjects_for_students, get_exam_students
		students = list(get_exam_students(self.exam))
		subjects, _is_filtered = get_exam_subjects_for_students(self.exam, students)
		self.assertNotIn(higher_math, subjects)

	def test_subject_scope_is_the_union_of_admission_assignments(self):
		from .result_utils import get_exam_subjects_for_students, get_exam_students
		students = list(get_exam_students(self.exam))
		subjects, is_filtered = get_exam_subjects_for_students(self.exam, students)
		self.assertTrue(is_filtered)
		self.assertEqual({subject.pk for subject in subjects}, {
			self.bangla.pk, self.ict.pk, self.agriculture.pk,
		})

	def test_unselected_optional_subject_is_blank_and_not_counted(self):
		from .result_utils import build_exam_results
		ExamMark.objects.create(exam=self.exam, student=self.ict_student, subject=self.bangla, marks_obtained=80)
		ExamMark.objects.create(exam=self.exam, student=self.ict_student, subject=self.ict, marks_obtained=70)
		ExamMark.objects.create(exam=self.exam, student=self.agriculture_student, subject=self.bangla, marks_obtained=80)
		ExamMark.objects.create(exam=self.exam, student=self.agriculture_student, subject=self.agriculture, marks_obtained=60)
		_columns, results = build_exam_results(self.exam)
		ict_result = next(r for r in results if r['student'].pk == self.ict_student.pk)
		agriculture_row = next(
			row for row in ict_result['subject_results']
			if row.get('subject') == self.agriculture
		)
		self.assertTrue(agriculture_row['subject_unassigned'])
		self.assertEqual(ict_result['total_full'], 200)
		self.assertEqual(ict_result['total_obtained'], 150)
		self.assertEqual(ict_result['status'], 'Pass')
		self.assertEqual(ict_result['absent_subject_count'], 0)

	def test_marks_input_is_enabled_only_for_assigned_student(self):
		response = self.client.get(reverse('enter_marks', args=[self.exam.pk, self.ict.pk]))
		self.assertEqual(response.status_code, 200)
		content = response.content.decode()
		self.assertIn(f'name="marks_{self.ict_student.pk}"', content)
		self.assertNotIn(f'name="marks_{self.agriculture_student.pk}"', content)

	@skipUnless(Workbook, 'openpyxl is required for Excel import tests')
	def test_excel_import_skips_a_student_who_did_not_choose_the_subject(self):
		from .marks_import import parse_subject_marks_workbook
		from .result_utils import get_exam_students
		workbook = Workbook()
		sheet = workbook.active
		sheet.append(['Roll', 'ID', 'Name', 'Marks'])
		sheet.append([1, self.ict_student.student_id, self.ict_student.name, 70])
		sheet.append([2, self.agriculture_student.student_id, self.agriculture_student.name, 70])
		validated, skipped, errors = parse_subject_marks_workbook(
			workbook, self.exam, self.ict, list(get_exam_students(self.exam)),
		)
		self.assertFalse(errors)
		self.assertEqual(skipped, 1)
		self.assertEqual([student.pk for student, _defaults in validated], [self.ict_student.pk])

	def test_summary_total_has_no_full_marks_denominator_and_actions_are_print_hidden(self):
		ExamMark.objects.create(exam=self.exam, student=self.ict_student, subject=self.bangla, marks_obtained=80)
		ExamMark.objects.create(exam=self.exam, student=self.ict_student, subject=self.ict, marks_obtained=70)
		response = self.client.get(reverse('exam_result_summary', args=[self.exam.pk]))
		self.assertContains(response, '150')
		self.assertNotContains(response, '150 / 300')
		self.assertContains(response, 'class="d-print-none"')
		self.assertContains(response, 'A4 landscape')


class ResultCountingRulesTests(TestCase):
	"""A result is Pass only when every subject is passed individually. A
	failed result gets no percentage and no position — they are not counted."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Counting School', classes='6')
		self.user = get_user_model().objects.create_superuser(username='counting-admin', password='password')
		self.client.force_login(self.user)
		from .models import SubjectRequirement
		self.bangla = Subject.objects.create(code='CBN', name='Bangla', full_marks=100)
		self.math = Subject.objects.create(code='CMT', name='Math', full_marks=100)
		for subject in (self.bangla, self.math):
			SubjectRequirement.objects.create(
				institution=self.institution, admission_class='6',
				subject=subject, requirement_type='MANDATORY',
			)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='6', session='2026', is_published=True,
		)
		self.passer = Student.objects.create(
			institution=self.institution, student_id='C001', name='All Pass',
			admission_class='6', section='A', roll_no=1, admission_year=2026,
		)
		self.failer = Student.objects.create(
			institution=self.institution, student_id='C002', name='One Fail',
			admission_class='6', section='A', roll_no=2, admission_year=2026,
		)

	def _result_for(self, student):
		from .result_utils import build_exam_results
		_, results = build_exam_results(self.exam)
		return next(r for r in results if r['student'].pk == student.pk)

	def test_pass_only_when_every_subject_is_passed_individually(self):
		ExamMark.objects.create(exam=self.exam, student=self.passer, subject=self.bangla, marks_obtained=72)
		ExamMark.objects.create(exam=self.exam, student=self.passer, subject=self.math, marks_obtained=65)
		ExamMark.objects.create(exam=self.exam, student=self.failer, subject=self.bangla, marks_obtained=78)
		ExamMark.objects.create(exam=self.exam, student=self.failer, subject=self.math, marks_obtained=30)

		passed = self._result_for(self.passer)
		self.assertEqual(passed['status'], 'Pass')
		self.assertEqual(passed['percentage'], 68.5)
		self.assertEqual(passed['position'], 1)

		failed = self._result_for(self.failer)
		self.assertEqual(failed['status'], 'Fail')
		self.assertEqual(str(failed['gpa']), '0.00')
		self.assertIsNone(failed['percentage'])
		self.assertIsNone(failed['position'])

	def test_summary_page_shows_a_dash_for_a_failed_percentage(self):
		ExamMark.objects.create(exam=self.exam, student=self.failer, subject=self.bangla, marks_obtained=78)
		ExamMark.objects.create(exam=self.exam, student=self.failer, subject=self.math, marks_obtained=30)
		response = self.client.get(reverse('exam_result_summary', args=[self.exam.pk]))
		self.assertContains(response, 'Fail')
		self.assertContains(response, '&mdash;')

	def test_class_zero_padding_still_finds_assigned_subjects(self):
		from .result_utils import get_exam_subjects
		# Requirements stored as '06' must match an exam stored as '6'.
		from .models import SubjectRequirement
		SubjectRequirement.objects.all().delete()
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='06',
			subject=self.bangla, requirement_type='MANDATORY',
		)
		subjects, is_filtered = get_exam_subjects(self.exam)
		self.assertTrue(is_filtered)
		self.assertIn(self.bangla, list(subjects))


class ReligionFormFieldTests(TestCase):
	"""The religion field is the Islam/Hindu dropdown, and legacy free-text
	values are normalised rather than rejected."""

	def test_religion_is_a_dropdown_with_the_two_papers(self):
		form = StudentForm()
		html = str(form['religion'])
		self.assertIn('<option value="Islam"', html)
		self.assertIn('<option value="Hindu"', html)
		self.assertIn('selected', html)  # Islam selected by default

	def test_legacy_value_is_normalised_to_the_nearest_paper(self):
		institution = Institution.objects.create(name='Form School', classes='6')
		student = Student.objects.create(
			institution=institution, student_id='F001', name='Legacy Kid',
			admission_class='6', admission_year=2026, religion='মুসলিম',
		)
		form = StudentForm(instance=student)
		self.assertIn('<option value="Islam" selected', str(form['religion']))

		student.religion = 'hinduism'
		student.save(update_fields=['religion'])
		form = StudentForm(instance=student)
		self.assertIn('<option value="Hindu" selected', str(form['religion']))

	def test_submitted_religion_is_saved_normalised(self):
		institution = Institution.objects.create(name='Form School 2', classes='6')
		user = get_user_model().objects.create_superuser(username='form-admin', password='password')
		self.client.force_login(user)
		self.client.post(reverse('add_student'), {
			'institution': institution.pk,
			'name': 'New Kid',
			'admission_class': '6',
			'section': 'A',
			'admission_year': '2026',
			'religion': 'Hindu',
			'status': 'ACTIVE',
		})
		student = Student.objects.get(name='New Kid')
		self.assertEqual(student.religion, 'Hindu')

	def test_get_applicable_subjects_uses_the_students_own_religion_paper_only(self):
		from .views import get_applicable_subjects
		from .models import SubjectRequirement
		institution = Institution.objects.create(name='Applicable School', classes='6')
		islam = Subject.objects.create(
			code='APISL', name='Islam & Moral Education', full_marks=100, category='RELIGION')
		hindu = Subject.objects.create(
			code='APHIN', name='Hindu Religion & Moral Education', full_marks=100, category='RELIGION')
		christian = Subject.objects.create(
			code='APCHR', name='Christian Religion & Moral Education', full_marks=100, category='RELIGION')
		for subject, religion in (
			(islam, 'Islam'), (hindu, 'Hindu'), (christian, 'Christian'),
		):
			SubjectRequirement.objects.create(
				institution=institution, admission_class='6', subject=subject,
				requirement_type='CONDITIONAL', condition_religion=religion,
			)

		for religion, expected in (
			('Hindu', hindu), ('', islam), ('Muslim', islam), ('islam', islam),
		):
			data = get_applicable_subjects(institution, '6', religion=religion)
			names = [item['name'] for item in data['conditional']]
			self.assertEqual(names, [expected.name], f"religion={religion!r}")

	def test_get_applicable_subjects_adds_no_paper_when_the_own_one_is_missing(self):
		from .views import get_applicable_subjects
		from .models import SubjectRequirement
		institution = Institution.objects.create(name='Applicable School 2', classes='6')
		islam = Subject.objects.create(
			code='APIS2', name='Islam & Moral Education', full_marks=100, category='RELIGION')
		SubjectRequirement.objects.create(
			institution=institution, admission_class='6', subject=islam,
			requirement_type='CONDITIONAL', condition_religion='Islam',
		)
		# A Hindu student with no Hindu paper assigned gets NO religion paper —
		# Islam is never substituted. The result sheet shows a REL dash.
		data = get_applicable_subjects(institution, '6', religion='Hindu')
		self.assertEqual(data['conditional'], [])
		# Muslim students still get the Islam paper.
		data = get_applicable_subjects(institution, '6', religion='Islam')
		self.assertEqual([item['name'] for item in data['conditional']], [islam.name])


class NoSubjectsAssignedTests(TestCase):
	"""The old 'nothing assigned -> show every subject' fallback is gone.
	get_exam_subjects returns an empty list for an unassigned class, and the
	Enter Marks, Excel Import and Result Sheet pages show a clear message
	pointing at the Subject Assignments page instead."""

	NO_SUBJECTS_MSG = (
		"No subjects are assigned to Class 6 yet — assign them from the "
		"Subject Assignments page first."
	)

	def setUp(self):
		self.institution = Institution.objects.create(name='Empty School', classes='6')
		self.user = get_user_model().objects.create_superuser(username='empty-admin', password='password')
		self.client.force_login(self.user)
		self.subject = Subject.objects.create(code='STR', name='Stray Subject', full_marks=100)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='6', session='2026', is_published=True,
		)
		self.student = Student.objects.create(
			institution=self.institution, student_id='E001', name='Empty Class Kid',
			admission_class='6', section='A', roll_no=1, admission_year=2026,
		)

	def test_get_exam_subjects_does_not_fall_back_to_every_subject(self):
		from .result_utils import get_exam_subjects
		# Even though the Subject master list holds a subject, an unassigned
		# class gets an empty list, not every subject in the system.
		self.assertGreater(Subject.objects.count(), 0)
		subjects, is_filtered = get_exam_subjects(self.exam)
		self.assertEqual(list(subjects), [])
		self.assertFalse(is_filtered)

	def test_get_exam_subjects_returns_assigned_subjects_once_configured(self):
		from .result_utils import get_exam_subjects
		from .models import SubjectRequirement
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='6', subject=self.subject,
			requirement_type='MANDATORY',
		)
		subjects, is_filtered = get_exam_subjects(self.exam)
		self.assertEqual(list(subjects), [self.subject])
		self.assertTrue(is_filtered)

	def test_select_marks_subject_page_shows_the_message(self):
		response = self.client.get(reverse('select_marks_subject', args=[self.exam.pk]))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'No subjects are assigned to Class 6')
		self.assertContains(response, 'Subject Assignments')
		# No subject picker is offered for an unassigned class.
		self.assertNotContains(response, 'name="subject"')

	def test_import_page_shows_the_message(self):
		response = self.client.get(reverse('import_exam_marks', args=[self.exam.pk]))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'No subjects are assigned to Class 6')
		self.assertContains(response, 'Subject Assignments')
		self.assertNotContains(response, 'Upload and Import')

	def test_result_sheet_shows_the_message(self):
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'No subjects are assigned to Class 6')
		self.assertContains(response, 'Subject Assignments')

	def test_enter_marks_for_unassigned_subject_bounces_with_the_message(self):
		# Editing the URL straight to an unassigned subject does not work:
		# the page redirects back with the no-subjects message (no fallback to
		# unfiltered subject entry).
		response = self.client.get(
			reverse('enter_marks', args=[self.exam.pk, self.subject.pk]),
		)
		self.assertRedirects(response, reverse('select_marks_subject', args=[self.exam.pk]))
		followed = self.client.get(reverse('select_marks_subject', args=[self.exam.pk]))
		self.assertContains(followed, 'No subjects are assigned to Class 6')

	def test_publishing_and_results_do_not_fabricate_subjects(self):
		# Stray marks in an unassigned class cannot invent columns either.
		ExamMark.objects.create(
			exam=self.exam, student=self.student, subject=self.subject, marks_obtained=80,
		)
		from .result_utils import build_exam_results
		columns, results = build_exam_results(self.exam)
		self.assertEqual(list(columns), [])
		# The stray mark is not printed as a column...
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		self.assertContains(response, 'No subjects are assigned to Class 6')
		# ...and it does not grade the student a pass.
		result = next(r for r in results if r['student'].pk == self.student.pk)
		self.assertEqual(result['status'], 'No Marks')

	def test_message_helper_text(self):
		from .result_utils import no_subjects_assigned_message
		self.assertEqual(no_subjects_assigned_message(self.exam), self.NO_SUBJECTS_MSG)


class ResultSheetCodeHeaderTests(TestCase):
	"""The result sheet header shows subject codes (BAN1, ENG1, REL…) and the
	full names live in a 'Subject codes' legend under the table; the student
	name column is widened."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Header School', classes='6')
		self.user = get_user_model().objects.create_superuser(username='header-admin', password='password')
		self.client.force_login(self.user)
		from .models import SubjectRequirement
		self.bangla = Subject.objects.create(code='BAN1', name='Bangla 1st Paper', full_marks=100)
		self.english = Subject.objects.create(code='ENG1', name='English 1st Paper', full_marks=100)
		for subject in (self.bangla, self.english):
			SubjectRequirement.objects.create(
				institution=self.institution, admission_class='6', subject=subject,
				requirement_type='MANDATORY',
			)
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='6', session='2026', is_published=True,
		)
		self.student = Student.objects.create(
			institution=self.institution, student_id='H001', name='Header Kid',
			admission_class='6', section='A', roll_no=1, admission_year=2026,
		)
		ExamMark.objects.create(exam=self.exam, student=self.student, subject=self.bangla, marks_obtained=80)
		ExamMark.objects.create(exam=self.exam, student=self.student, subject=self.english, marks_obtained=70)

	def test_header_uses_codes_and_legend_maps_them_to_names(self):
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		content = response.content.decode()
		# Codes in the header row.
		self.assertIn('<th>BAN1<br>', content)
		self.assertIn('<th>ENG1<br>', content)
		# Full names do not appear as column headers.
		self.assertNotIn('<th>Bangla 1st Paper', content)
		# Legend under the table.
		self.assertIn('Subject codes', content)
		self.assertIn('<span class="code">BAN1</span> = Bangla 1st Paper', content)
		self.assertIn('<span class="code">ENG1</span> = English 1st Paper', content)

	def test_student_name_column_is_wide(self):
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]))
		self.assertContains(response, 'student-col')
		self.assertContains(response, 'min-width: 200px')


class SubjectAssignmentsConditionalNoteTests(TestCase):
	"""The Conditional section of Subject Assignments explains the two religion
	rows and the single merged Religion column on the result sheet."""

	def setUp(self):
		self.institution = Institution.objects.create(name='Conditional School', classes='6')
		self.user = get_user_model().objects.create_superuser(username='cond-admin', password='password')
		self.client.force_login(self.user)
		from .models import SubjectRequirement
		self.islam = Subject.objects.create(
			code='CISL', name='Islam & Moral Education', full_marks=100, category='RELIGION')
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='6', subject=self.islam,
			requirement_type='CONDITIONAL', condition_religion='Islam',
		)

	def test_conditional_section_explains_the_two_religion_rows(self):
		response = self.client.get(
			reverse('subject_requirement_list')
			+ f'?institution={self.institution.pk}&admission_class=6',
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'two religion rows')
		self.assertContains(response, 'one Religion column')


class ResultSheetSubjectAssignmentConnectionTests(TestCase):
	def setUp(self):
		from .models import SubjectRequirement
		self.institution = Institution.objects.create(name='Group Link School', classes='9')
		self.user = get_user_model().objects.create_superuser(username='group-link-admin', password='password')
		self.client.force_login(self.user)
		self.business = Subject.objects.create(code='BUS101', name='Business Studies', full_marks=100)
		self.science = Subject.objects.create(code='SCI101', name='Science', full_marks=100)
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', group='BUS',
			subject=self.business, requirement_type='MANDATORY',
		)
		SubjectRequirement.objects.create(
			institution=self.institution, admission_class='9', group='SCI',
			subject=self.science, requirement_type='MANDATORY',
		)
		self.bus_student = Student.objects.create(
			institution=self.institution, student_id='BUS001', name='Business Kid',
			admission_class='9', group='BUS', section='A', roll_no=1, admission_year=2026,
		)
		self.sci_student = Student.objects.create(
			institution=self.institution, student_id='SCI001', name='Science Kid',
			admission_class='9', group='SCI', section='A', roll_no=2, admission_year=2026,
		)
		# Older exams may be created for the whole class; the result sheet must be
		# able to follow the group selected from Subject Assignments.
		self.exam = Exam.objects.create(
			name='Second Term Examination-2026', exam_type='SECOND_TERM',
			institution=self.institution, admission_class='9', session='2026', is_published=True,
		)
		ExamMark.objects.create(exam=self.exam, student=self.bus_student, subject=self.business, marks_obtained=80)
		ExamMark.objects.create(exam=self.exam, student=self.sci_student, subject=self.science, marks_obtained=70)

	def test_subject_assignment_filter_links_to_matching_group_result_sheet(self):
		response = self.client.get(reverse('subject_requirement_list'), {
			'institution': self.institution.pk,
			'admission_class': '9',
			'group': 'BUS',
		})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Published Result Sheet')
		self.assertContains(response, f"{reverse('result_sheet', args=[self.exam.pk])}?group=BUS")

	def test_result_sheet_group_query_filters_students_and_links_back_to_assignments(self):
		response = self.client.get(reverse('result_sheet', args=[self.exam.pk]), {'group': 'BUS'})
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, 'Business Kid')
		self.assertNotContains(response, 'Science Kid')
		self.assertContains(response, f"{reverse('subject_requirement_list')}?institution={self.institution.pk}&amp;admission_class=9&amp;group=BUS")

	def test_grouped_exam_assignment_url_keeps_exam_group(self):
		from .views import subject_assignments_url
		exam = Exam.objects.create(
			name='Grouped Exam', exam_type='SECOND_TERM', institution=self.institution,
			admission_class='9', group='BUS', session='2026', is_published=True,
		)
		self.assertEqual(
			subject_assignments_url(exam),
			f"{reverse('subject_requirement_list')}?institution={self.institution.pk}&admission_class=9&group=BUS",
		)
