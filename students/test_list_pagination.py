"""OF-02: a hard 100-row UI cap, adjustable sizes, stable pages and full exports.

Only synthetic data in Django's disposable test database. Every list surface
is exercised with >100 rows, including grouped/multi-table reports and admin
reference lists. Print modes run through the same permission/scope guards.
"""
from datetime import date
from html.parser import HTMLParser
from io import BytesIO
from urllib.parse import parse_qs, urlsplit

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import RequestFactory, SimpleTestCase, TestCase
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    AdmissionApplication, AttendanceRecord, AuditLog, Certificate, Employee,
    EmployeeStatusLog, Exam, ExamMark, Fee, Institution, InstitutionAccess,
    MoneyReceipt, PromotionBatch, SeatPlan, Student, Subject, SubjectRequirement,
    Voucher, SalarySheet,
)
from .pagination import get_page_size, paginate_list
from .templatetags.student_filters import retained_list_filters


class PageMarkup(HTMLParser):
    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.links = []
        self.forms = []
        self._form = None
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'a':
            self.links.append(attrs)
        elif tag == 'form':
            self._form = {'attrs': attrs, 'fields': []}
            self.forms.append(self._form)
        elif tag == 'input' and self._form is not None:
            self._form['fields'].append(attrs)

    def handle_endtag(self, tag):
        if tag == 'form':
            self._form = None


class PaginationHelperTests(SimpleTestCase):
    def test_default_invalid_and_oversized_values_never_exceed_100(self):
        for raw in (None, '', 'oops', '0', '-1', '1.5', '9' * 5000, '101', '999999'):
            with self.subTest(raw=str(raw)[:20]):
                self.assertEqual(get_page_size(raw), 100)
        for size in (1, 7, 25, 50, 99, 100):
            self.assertEqual(get_page_size(str(size)), size)

    def test_invalid_pages_use_get_page_and_keep_the_cap(self):
        for raw, expected_page in [('oops', 1), ('999', 5), ('0', 5), ('-1', 5)]:
            request = RequestFactory().get('/', {'page': raw, 'per_page': 25})
            context = paginate_list(request, list(range(105)))
            self.assertEqual(context['page_obj'].number, expected_page)
            self.assertLessEqual(len(context['page_rows']), 25)

    def test_print_is_an_explicit_opt_in_not_a_generic_cap_bypass(self):
        request = RequestFactory().get('/', {'per_page': 1, 'print': '1', 'page': 99})
        rows = list(range(105))
        self.assertEqual(len(paginate_list(request, rows)['page_rows']), 1)
        printable = paginate_list(request, rows, allow_full_print=True)
        self.assertEqual(printable['page_rows'], rows)
        self.assertTrue(printable['is_print_view'])

    def test_links_and_forms_preserve_repeated_encoded_filters_and_reset_page(self):
        filters = {'q': 'A&B + "<script>"', 'institution': '42', 'sort': '-name',
                   'exams': ['7', '9'], 'per_page': '25', 'page': '2'}
        request = RequestFactory().get('/students/', filters)
        context = {'request': request, **paginate_list(request, list(range(105)))}
        html = render_to_string('students/_pagination.html', context)
        markup = PageMarkup(html)
        for label, page in [('First page', '1'), ('Previous page', '1'),
                            ('Next page', '3'), ('Last page', '5')]:
            link = next(link for link in markup.links if link.get('aria-label') == label)
            query = parse_qs(urlsplit(link['href']).query)
            self.assertEqual(query['page'], [page])
            for key in ('q', 'institution', 'sort', 'exams', 'per_page'):
                expected = filters[key] if isinstance(filters[key], list) else [filters[key]]
                self.assertEqual(query[key], expected)
        form = next(form for form in markup.forms if 'data-page-size-form' in form['attrs'])
        names = [field.get('name') for field in form['fields']]
        self.assertNotIn('page', names)
        self.assertEqual(names.count('exams'), 2)
        self.assertEqual(names.count('per_page'), 1)
        self.assertNotIn('<script>', html)

    def test_filter_fields_exclude_the_callers_inputs_and_keep_duplicates(self):
        request = RequestFactory().get('/', {'q': 'a&b', 'institution': '1',
                                             'exams': ['7', '9'], 'page': 2,
                                             'per_page': 25, 'print': 1})
        fields = retained_list_filters({'request': request}, 'q institution')
        self.assertEqual(fields, [('exams', '7'), ('exams', '9')])

    def test_empty_and_single_page_lists_still_offer_size_controls_and_counts(self):
        for rows in ([], [1]):
            request = RequestFactory().get('/')
            html = render_to_string('students/_pagination.html', {
                'request': request, **paginate_list(request, rows),
            })
            self.assertIn('data-page-size-form', html)
            self.assertIn('Page 1 of 1', html)
            self.assertIn(f'{len(rows)} total', html)
            self.assertIn('max="100"', html)


class ListPaginationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_superuser('page-admin', 'page@example.invalid', 'synthetic-only')
        cls.institution = Institution.objects.create(name='Pagination School', classes='6,9')
        cls.today = date.today()
        cls.now = timezone.now()
        cls.students = Student.objects.bulk_create([
            Student(institution=cls.institution, student_id=f'PG{i:03d}', name='Page child & +',
                    admission_class='6', section='A', roll_no=1,
                    guardian_contact_no='01800000000', admission_year=2026)
            for i in range(105)
        ])
        cls.archived = Student.objects.bulk_create([
            Student(institution=cls.institution, student_id=f'AR{i:03d}', name='Archived page child',
                    admission_class='6', section='A', roll_no=1, is_archived=True,
                    archived_at=cls.now, guardian_contact_no='01800000000', admission_year=2026)
            for i in range(105)
        ])
        # 105 distinct reference/aggregate rows in addition to Section A.
        Student.objects.bulk_create([
            Student(institution=cls.institution, student_id=f'SC{i:03d}', name='Section reference',
                    admission_class='6', section=f'S{i:03d}', roll_no=1,
                    guardian_contact_no='01800000000', admission_year=2026)
            for i in range(105)
        ])
        cls.employees = Employee.objects.bulk_create([
            Employee(institution=cls.institution, name='Page employee', designation='Teacher',
                     join_date=cls.today, status='ACTIVE') for i in range(105)
        ])
        cls.subject = Subject.objects.create(code='PGMARK', name='Pagination marks', full_marks=100)
        SubjectRequirement.objects.create(institution=cls.institution, admission_class='6',
                                          subject=cls.subject, requirement_type='MANDATORY')
        subjects = Subject.objects.bulk_create([
            Subject(code=f'PGREF{i:03d}', name='Pagination reference', full_marks=100)
            for i in range(105)
        ])
        SubjectRequirement.objects.bulk_create([
            SubjectRequirement(institution=cls.institution, admission_class='9', group='SCI',
                               subject=subject, requirement_type='MANDATORY') for subject in subjects
        ])
        cls.exams = Exam.objects.bulk_create([
            Exam(institution=cls.institution, name='Pagination Exam', admission_class='6',
                 section='A', exam_type='FIRST_TERM', session='2026', exam_date=cls.today,
                 is_published=True) for i in range(105)
        ])
        ExamMark.objects.bulk_create([
            ExamMark(exam=exam, student=student, subject=cls.subject,
                     marks_obtained=80 if i < 100 else 20)
            for exam in cls.exams[:2] for i, student in enumerate(cls.students)
        ])
        AdmissionApplication.objects.bulk_create([
            AdmissionApplication(institution=cls.institution, application_number=f'PGAPP{status}{i:03d}',
                                 applicant_name='Page applicant', guardian_name='Synthetic guardian',
                                 guardian_contact_no='01800000000', requested_class='6', session='2026',
                                 status=status)
            for status in ('SUBMITTED', 'ACCOUNT_PENDING') for i in range(105)
        ])
        references = Institution.objects.bulk_create([
            Institution(name=f'Pagination reference {i:03d}', classes='6') for i in range(105)
        ])
        AdmissionApplication.objects.bulk_create([
            AdmissionApplication(institution=institution, application_number=f'PGREFAPP{i:03d}',
                                 applicant_name='Synthetic reference', guardian_name='Synthetic guardian',
                                 guardian_contact_no='01800000000', requested_class='6', session='2026',
                                 status='REJECTED') for i, institution in enumerate(references)
        ])
        AdmissionApplication.objects.update(submitted_at=cls.now)
        MoneyReceipt.objects.bulk_create([
            MoneyReceipt(student=cls.students[0], receipt_no=f'PGRC{i:03d}', purpose='Synthetic fee',
                         amount=10, date=cls.today) for i in range(105)
        ])
        Voucher.objects.bulk_create([
            Voucher(institution=cls.institution, purpose='Synthetic voucher', amount=10,
                    date=cls.today) for i in range(105)
        ])
        SalarySheet.objects.bulk_create([
            SalarySheet(employee=cls.employees[0], month=f'Synthetic {i:03d}', amount=10,
                        date=cls.today) for i in range(105)
        ])
        Fee.objects.bulk_create([
            Fee(institution=cls.institution, admission_class='6', purpose=f'Synthetic fee {i:03d}', amount=10)
            for i in range(105)
        ])
        AuditLog.objects.bulk_create([
            AuditLog(institution=cls.institution, action='pagination_fixture', model_name='students.Student',
                     object_id=str(i), object_repr=f'Synthetic {i:03d}') for i in range(105)
        ])
        AuditLog.objects.update(timestamp=cls.now)
        PromotionBatch.objects.bulk_create([
            PromotionBatch(institution=cls.institution, session='2026', from_class='6', to_class='7',
                           from_section='A', to_section='B', actor=cls.user) for i in range(105)
        ])
        PromotionBatch.objects.update(created_at=cls.now)
        EmployeeStatusLog.objects.bulk_create([
            EmployeeStatusLog(employee=cls.employees[0], old_status='ACTIVE', new_status='INACTIVE',
                              reason='Synthetic history') for i in range(105)
        ])
        EmployeeStatusLog.objects.update(changed_at=cls.now)
        Certificate.objects.bulk_create([
            Certificate(student=cls.students[0], certificate_type='CHARACTER',
                        certificate_number=f'PGCERT{i:03d}', purpose='Synthetic certificate') for i in range(105)
        ])
        SeatPlan.objects.bulk_create([
            SeatPlan(exam=cls.exams[0], student=student, room_name='Whole Room',
                     room_type='CLASSROOM', seat_no=i + 1) for i, student in enumerate(cls.students)
        ] + [
            SeatPlan(exam=cls.exams[1], student=student, room_name=f'Room {i:03d}',
                     room_type='CLASSROOM', seat_no=1) for i, student in enumerate(cls.students)
        ])
        AttendanceRecord.objects.bulk_create([
            AttendanceRecord(institution=cls.institution, student=student, date=cls.today, status='P')
            for student in cls.students
        ] + [
            AttendanceRecord(institution=cls.institution, employee=employee, date=cls.today, status='A')
            for employee in cls.employees
        ])

    def setUp(self):
        self.client.force_login(self.user)

    def url_and_params(self, name):
        pk = self.students[0].pk if name in ('student_exams', 'certificate_list') else self.exams[0].pk
        if name in ('employee_status_history', 'employee_detail'):
            pk = self.employees[0].pk
        if name == 'seat_plan_list':
            pk = self.exams[1].pk
        args = [pk] if name in ('student_exams', 'certificate_list', 'employee_status_history',
                               'seat_plan_list', 'result_sheet', 'exam_result_summary', 'full_rank_list', 'employee_detail') else []
        if name == 'view_seat_plan_room':
            args = [self.exams[0].pk, 'Whole Room']
        params = {}
        if name == 'student_list':
            params = {'institution': self.institution.pk, 'section': 'A', 'q': 'Page child & +'}
        if name == 'employee_list':
            params = {'institution': self.institution.pk, 'status': 'ACTIVE'}
        if name == 'subject_requirement_list':
            params = {'institution': self.institution.pk, 'admission_class': '9', 'group': 'SCI', 'q': 'Pagination reference'}
        if name == 'admission_application_list':
            params = {'status': 'SUBMITTED'}
        if name == 'accounts_admission_queue':
            params = {'admission_class': '6'}
        if name.startswith('result_analysis_') or name == 'section_arrangement':
            params = {'exam': self.exams[0].pk}
        if name == 'result_analysis_multi_term':
            params = {'exams': [exam.pk for exam in self.exams[:2]]}
        if name == 'result_analysis_subject_fail':
            ExamMark.objects.filter(exam=self.exams[0]).update(marks_obtained=20)
        return reverse(name, args=args), params

    def assert_list_surface(self, name, count):
        url, params = self.url_and_params(name)
        first = self.client.get(url, params)
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.context['paginator'].count, count)
        self.assertEqual(len(first.context['page_rows']), 100)
        self.assertContains(first, 'data-page-size-form')
        self.assertContains(first, f'{count} total')
        second = self.client.get(url, {**params, 'page': 2})
        self.assertEqual(second.status_code, 200)
        self.assertEqual(len(second.context['page_rows']), min(100, count - 100))
        # Walk all pages to prove each underlying row occurs exactly once.
        def identities(rows):
            keys = []
            for row in rows:
                if hasattr(row, 'pk'):
                    keys.append(row.pk)
                elif isinstance(row, tuple):
                    keys.append(row[0])  # room name
                elif 'subject' in row:
                    keys.append((row['subject'].pk, row['student'].pk))
                elif 'student' in row:
                    keys.append(('student', row['student'].pk))
                elif 'employee' in row:
                    keys.append(('employee', row['employee'].pk))
                elif 'admission_class' in row:
                    keys.append((row['admission_class'], row['section']))
                else:
                    keys.append(row['institution'])
            return keys
        seen = identities(first.context['page_rows']) + identities(second.context['page_rows'])
        for page in range(3, first.context['paginator'].num_pages + 1):
            response = self.client.get(url, {**params, 'page': page})
            seen += identities(response.context['page_rows'])
        self.assertEqual(len(seen), count)
        self.assertEqual(len(set(seen)), count)
        for extra, expected_size, expected_number in [
            ({'per_page': 37, 'page': 2}, 37, 2),
            ({'per_page': 1000}, 100, 1),
            ({'per_page': 'oops'}, 100, 1),
            ({'page': 'oops'}, 100, 1),
            ({'page': 999999}, count - 100 * ((count - 1) // 100), (count + 99) // 100),
        ]:
            with self.subTest(options=extra):
                response = self.client.get(url, {**params, **extra})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(len(response.context['page_rows']), expected_size)
                self.assertEqual(response.context['page_obj'].number, expected_number)
        if name == 'subject_requirement_list':
            self.assertEqual(sum(map(len, first.context['grouped_requirements'].values())), 100)
        if name == 'full_rank_list':
            self.assertEqual(len(first.context['results']) + len(first.context['unranked_results']), 100)
        if name == 'attendance_summary':
            self.assertEqual(len(first.context['student_summary']) + len(first.context['employee_summary']), 100)
        if name == 'result_analysis_subject_fail':
            self.assertEqual(sum(len(bucket['rows']) for bucket in first.context['subjects_with_fails']), 100)
            self.assertEqual(first.context['subjects_with_fails'][0]['count'], 105)

    def test_size_one_and_filters_survive_links_and_filter_forms(self):
        params = {'institution': self.institution.pk, 'status': 'ACTIVE', 'per_page': 1,
                  'sort': '-name', 'tag': ['a&b', 'c+d']}
        response = self.client.get(reverse('employee_list'), params)
        self.assertEqual(len(response.context['employees']), 1)
        markup = PageMarkup(response.content.decode())
        link = next(link for link in markup.links if link.get('aria-label') == 'Next page')
        query = parse_qs(urlsplit(link['href']).query)
        self.assertEqual(query['page'], ['2'])
        self.assertEqual(query['tag'], ['a&b', 'c+d'])
        self.assertEqual(query['per_page'], ['1'])
        self.assertEqual(query['status'], ['ACTIVE'])
        form = next(form for form in markup.forms if form['attrs'].get('method', '').lower() == 'get'
                    and 'data-page-size-form' not in form['attrs'])
        hidden = [(f.get('name'), f.get('value')) for f in form['fields'] if f.get('type') == 'hidden']
        self.assertIn(('per_page', '1'), hidden)
        self.assertIn(('sort', '-name'), hidden)
        self.assertIn(('tag', 'a&b'), hidden)
        self.assertNotIn('page', [name for name, value in hidden])
        self.assertContains(response, 'Total: 105')

    def test_larger_pages_do_not_add_a_related_object_query_for_each_row(self):
        EmployeeStatusLog.objects.update(changed_by=self.user)
        for name in ('student_list', 'employee_list', 'exam_list', 'employee_detail'):
            with self.subTest(view=name):
                url, params = self.url_and_params(name)
                with CaptureQueriesContext(connection) as small:
                    response = self.client.get(url, {**params, 'per_page': 1})
                    self.assertEqual(response.status_code, 200)
                with CaptureQueriesContext(connection) as large:
                    response = self.client.get(url, {**params, 'per_page': 100})
                    self.assertEqual(response.status_code, 200)
                self.assertLessEqual(len(large), len(small) + 2)
        # Aggregation visits 210 people, but their FKs are fetched in two joins,
        # not one query per person. Counts/rates are still full-cohort values.
        with CaptureQueriesContext(connection) as summary:
            response = self.client.get(reverse('attendance_summary'))
            self.assertEqual(response.status_code, 200)
        self.assertLessEqual(len(summary), 25)

    def test_result_and_report_totals_are_not_recomputed_on_a_page(self):
        response = self.client.get(reverse('result_sheet', args=[self.exams[0].pk]), {'per_page': 1, 'page': 105})
        self.assertEqual(len(response.context['results']), 1)
        self.assertEqual(response.context['results'][0]['status'], 'Fail')
        self.assertEqual(response.context['pass_rate'], 95)
        self.assertEqual(response.context['total_results'], 105)
        self.assertContains(response, 'id="passRate">95%')
        response = self.client.get(reverse('class_section_summary'), {'per_page': 1})
        self.assertEqual(response.context['grand_total'], 210)
        response = self.client.get(reverse('attendance_summary'), {'per_page': 1})
        self.assertEqual(response.context['total_records'], 210)
        self.assertEqual(response.context['student_summary_count'], 105)
        self.assertEqual(response.context['employee_summary_count'], 105)
        response = self.client.get(reverse('seat_plan_list', args=[self.exams[1].pk]), {'per_page': 1})
        self.assertEqual(response.context['seated_count'], 105)
        self.assertEqual(response.context['total_students'], 105)
        response = self.client.get(reverse('admission_funnel_report'), {'per_page': 1})
        self.assertEqual(response.context['total_applications'], 315)

    def test_print_all_reports_ignore_ui_page_but_preserve_filters(self):
        for name, count in [('student_list', 105), ('employee_list', 105), ('result_sheet', 105),
                            ('exam_result_summary', 105), ('full_rank_list', 105),
                            ('view_seat_plan_room', 105), ('class_section_summary', 106),
                            ('attendance_summary', 210), ('admission_funnel_report', 106),
                            ('result_analysis_multi_term', 105), ('section_arrangement', 105),
                            ('result_analysis_subject_fail', 105)]:
            with self.subTest(view=name):
                url, params = self.url_and_params(name)
                response = self.client.get(url, {**params, 'print': 1, 'per_page': 1, 'page': 2})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.context['is_print_view'])
                self.assertEqual(len(response.context['page_rows']), count)
                self.assertNotContains(response, 'data-list-pagination')
                self.assertContains(response, 'Back to paginated list')

    def test_excel_exports_and_dedicated_prints_stay_full_data(self):
        expected_ids = {student.student_id for student in self.students}
        def workbook_values(response):
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(BytesIO(response.content))
            return {str(value) for sheet in workbook for row in sheet.iter_rows(values_only=True)
                    for value in row if value is not None}
        values = workbook_values(self.client.get(reverse('download_student_list'), {
            'institution': self.institution.pk, 'section': 'A', 'per_page': 1, 'page': 2,
        }))
        self.assertTrue(expected_ids.issubset(values))
        values = workbook_values(self.client.get(reverse('download_admission_sheet'), {
            'status': 'SUBMITTED', 'per_page': 1, 'page': 2,
        }))
        self.assertTrue({f'PGAPPSUBMITTED{i:03d}' for i in range(105)}.issubset(values))
        values = workbook_values(self.client.get(reverse('download_marks_import_template', args=[self.exams[0].pk]), {
            'subject': self.subject.pk, 'per_page': 1, 'page': 2,
        }))
        self.assertTrue(expected_ids.issubset(values))
        values = workbook_values(self.client.get(reverse('admission_funnel_export'), {'per_page': 1, 'page': 2}))
        self.assertTrue({f'Pagination reference {i:03d}' for i in range(105)}.issubset(values))
        response = self.client.get(reverse('signature_sheet', args=[self.exams[0].pk, 'Whole Room']), {'per_page': 1, 'page': 2})
        self.assertEqual(len(response.context['seats']), 105)
        response = self.client.get(reverse('result_analysis_result_cards'), {'exam': self.exams[0].pk, 'per_page': 1, 'page': 2})
        self.assertEqual(len(response.context['results']), 105)

    def test_arrangement_confirmation_uses_all_rows_not_only_visible_page(self):
        url = reverse('section_arrangement')
        params = {'exam': self.exams[0].pk, 'sections': 'A,B', 'per_page': 1, 'page': 105,
                  'institution': self.institution.pk, 'admission_class': '6'}
        preview = self.client.get(url, {**params, 'print': 1})
        expected = {row['student'].pk: row['proposed_section'] for row in preview.context['rows']}
        response = self.client.post(url + '?' + '&'.join(f'{k}={v}' for k, v in params.items()), {
            'exam': self.exams[0].pk, 'sections': 'A,B', 'action': 'confirm',
        })
        self.assertEqual(response.status_code, 302)
        actual = dict(Student.objects.filter(pk__in=expected).values_list('pk', 'section'))
        self.assertEqual(actual, expected)
        redirected = parse_qs(urlsplit(response.url).query)
        self.assertEqual(redirected['per_page'], ['1'])
        self.assertEqual(redirected['institution'], [str(self.institution.pk)])
        self.assertNotIn('page', redirected)
        self.assertEqual(Student.objects.filter(pk__in=expected, roll_no=1).count(), 105)
        self.assertEqual(AuditLog.objects.get(action='section_arrangement_confirmed').snapshot['student_count'], 105)

    def test_inline_employee_history_is_bounded_and_retains_the_active_tab(self):
        response = self.client.get(reverse('employee_detail', args=[self.employees[0].pk]))
        self.assertEqual(len(response.context['status_logs']), 100)
        response = self.client.get(reverse('employee_detail', args=[self.employees[0].pk]), {'page': 2, 'per_page': 7})
        self.assertEqual(len(response.context['status_logs']), 7)
        self.assertTrue(response.context['show_status_history'])
        self.assertContains(response, 'class="tab-pane fade show active" id="status"')
        self.assertContains(response, reverse('employee_status_history', args=[self.employees[0].pk]))

    def test_admin_reference_lists_are_selectable_and_show_all_cannot_bypass_cap(self):
        for model in ('institution', 'fee', 'employee', 'subjectrequirement'):
            with self.subTest(model=model):
                url = reverse(f'admin:students_{model}_changelist')
                response = self.client.get(url, {'per_page': 7, 'p': 2})
                self.assertEqual(response.status_code, 200)
                changelist = response.context['cl']
                self.assertEqual(changelist.list_per_page, 7)
                self.assertEqual(len(changelist.result_list), 7)
                self.assertEqual(changelist.page_num, 2)
                self.assertContains(response, 'data-page-size-form')
                for query in ({'all': '', 'per_page': 500}, {'per_page': 100, 'p': 'oops'},
                              {'per_page': 100, 'p': 99999}, {'per_page': -1, 'p': 0}):
                    response = self.client.get(url, query)
                    self.assertEqual(response.status_code, 200)
                    self.assertLessEqual(len(response.context['cl'].result_list), 100)
                    self.assertFalse(response.context['cl'].can_show_all)
                    self.assertNotContains(response, 'class="showall"')

    def test_admin_search_and_order_keep_the_chosen_size(self):
        response = self.client.get(reverse('admin:students_subjectrequirement_changelist'), {
            'q': 'Pagination reference', 'o': '1', 'per_page': 25,
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['cl'].result_count, 105)
        self.assertEqual(len(response.context['cl'].result_list), 25)
        self.assertContains(response, 'per_page=25')
        self.assertContains(response, 'q=Pagination+reference')

    def test_scoped_user_cannot_escape_via_size_page_institution_or_print(self):
        outsider = Institution.objects.create(name='Hidden institution', classes='6')
        Student.objects.create(institution=outsider, student_id='HIDDEN', name='Hidden student',
                               admission_class='6', section='A', guardian_contact_no='01800000000')
        AuditLog.objects.create(institution=outsider, action='hidden_audit')
        AuditLog.objects.create(action='legacy_system_audit')
        user = get_user_model().objects.create_user('page-clerk', password='synthetic-only')
        user.user_permissions.set(Permission.objects.filter(content_type__app_label='students'))
        InstitutionAccess.objects.create(user=user, institution=self.institution, department='Office')
        self.client.force_login(user)
        session = self.client.session
        session['selected_institution_id'] = outsider.pk
        session['selected_department'] = 'Office'
        session.save()
        for name in ('student_list', 'archived_students', 'employee_list', 'audit_log_list'):
            with self.subTest(view=name):
                response = self.client.get(reverse(name), {'institution': outsider.pk, 'per_page': 1,
                                                            'page': 2, 'print': 1})
                self.assertEqual(response.status_code, 200)
                for row in response.context['page_rows']:
                    self.assertEqual(row.institution_id, self.institution.pk)
                self.assertNotContains(response, 'Hidden student')
                self.assertNotContains(response, 'hidden_audit')
                self.assertNotContains(response, 'legacy_system_audit')
        response = self.client.get(reverse('student_detail', args=[Student.objects.get(student_id='HIDDEN').pk]), {'print': 1})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.client.get(reverse('admin:students_institution_changelist')).status_code, 302)


# Separate named tests make the inventory's coverage and failures visible in CI.
LIST_SURFACES = {
    'student_list': 105, 'archived_students': 105, 'employee_list': 105,
    'attendance_report': 210, 'admission_application_list': 105,
    'accounts_admission_queue': 105, 'money_receipt_list': 105, 'voucher_list': 105,
    'salary_sheet_list': 105, 'audit_log_list': 105, 'exam_list': 105,
    'subject_requirement_list': 105, 'student_exams': 105, 'certificate_list': 105,
    'student_promotion_history': 105, 'employee_status_history': 105, 'employee_detail': 105,
    'seat_plan_list': 105, 'view_seat_plan_room': 105, 'result_sheet': 105,
    'exam_result_summary': 105, 'full_rank_list': 105,
    'result_analysis_subject_fail': 105, 'result_analysis_multi_term': 105,
    'section_arrangement': 105, 'class_section_summary': 106,
    'attendance_summary': 210, 'admission_funnel_report': 106,
}
for _name, _count in LIST_SURFACES.items():
    def _test(self, name=_name, count=_count):
        self.assert_list_surface(name, count)
    _test.__name__ = f'test_{_name}_bounded_stable_pages'
    setattr(ListPaginationTests, _test.__name__, _test)
