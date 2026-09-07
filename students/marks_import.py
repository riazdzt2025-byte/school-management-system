"""Subject-wise Excel marks import — one file per subject, matching how
teachers already fill sheets (Roll, ID, Name, CQ, MCQ, PT).
"""
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile

from .models import ExamMark, MARK_PARTS, Student, normalize_class_label
from .result_utils import get_exam_students, get_subject_marks


EXCEL_PART_COLUMN = {
    'cq': 'CQ',
    'mcq': 'MCQ',
    'practical': 'PT',
    'weekly_test': 'WT',
}

PART_HEADER_ALIASES = {
    'cq': {'cq'},
    'mcq': {'mcq'},
    'practical': {'pt', 'practical', 'prac', 'practical marks'},
    'weekly_test': {'wt', 'weekly', 'weekly test', 'weekly_test'},
}

GROUP_SHEET_SHORT = {'SCI': 'SC', 'BUS': 'BUS', 'HUM': 'HUM'}


def _header_key(value):
    return re.sub(r'\s+', ' ', str(value or '').strip().lower())


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def marks_import_sheet_title(exam, subject):
    cls = normalize_class_label(exam.admission_class) or str(exam.admission_class)
    group_short = GROUP_SHEET_SHORT.get(exam.group or '', exam.group or '')
    title = f"{cls}{group_short} {subject.name}".strip()
    safe = ''.join(c for c in title if c not in r':\/?*[]')
    return (safe or 'Marks')[:31]


def marks_import_headers(marks_config):
    headers = ['Roll', 'ID', 'Name']
    if marks_config.parts:
        headers.extend(EXCEL_PART_COLUMN[part['key']] for part in marks_config.parts)
    else:
        headers.append('Marks')
    return headers


def build_subject_marks_workbook(exam, subject, group=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Protection
    from openpyxl.utils import get_column_letter

    marks_config = get_subject_marks(exam, subject)
    students = list(get_exam_students(exam, group=group))
    headers = marks_import_headers(marks_config)
    parts = marks_config.parts

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = marks_import_sheet_title(exam, subject)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for student in students:
        row = [student.roll_no, student.student_id, student.name]
        row.extend([None] * (len(headers) - 3))
        sheet.append(row)

    widths = [10, 16, 32] + [10] * (len(headers) - 3)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.protection.sheet = False
    for cell in sheet[1]:
        cell.protection = Protection(locked=True)

    guide = workbook.create_sheet('How to fill')
    guide.column_dimensions['A'].width = 100
    part_line = (
        ' + '.join(f"{part['label']} {part['max_marks']}" for part in parts)
        if parts else f"single total out of {marks_config.full_marks}"
    )
    for line in [
        f"{exam.name} — Class {exam.admission_class}"
        + (f" ({exam.section})" if exam.section else '')
        + (f" ({exam.get_group_display()} group)" if exam.group else ''),
        f"Subject: {subject.name} ({subject.code}) — {part_line}",
        '',
        'This file is for ONE subject only. The Physics teacher fills Physics;',
        'the Bangla teacher fills Bangla. Do not add other subjects to this sheet.',
        '',
        '1. Fill CQ / MCQ / PT (and WT if present). Do not change Roll, ID or Name.',
        '2. Leave a cell blank when the student did not sit that paper — blank is',
        '   skipped, while 0 is a real mark of zero.',
        '3. A mark above that part\'s maximum is rejected and the whole file is not imported.',
        f'4. {len(students)} student(s) are listed from the CURRENT class roll at download time.',
        '   If students join or leave, download this file again — do not reuse an old sheet.',
        '   Extra rows for students who have left are skipped on import.',
    ]:
        guide.append([line])

    workbook.active = 0
    return workbook


def workbook_as_upload(workbook, filename='marks.xlsx'):
    output = BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(
        filename, output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


def pick_marks_sheet(workbook):
    """The filled marks sheet, not the How to fill / Students helper tabs."""
    skipped = {'how to fill', 'students', 'subjects'}
    for sheet in workbook.worksheets:
        if (sheet.title or '').strip().lower() in skipped:
            continue
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = _column_map(header)
        if 'id' in columns or 'roll' in columns:
            return sheet
    return workbook.active


def parse_subject_marks_workbook(workbook, exam, subject, students):
    return parse_subject_marks_sheet(pick_marks_sheet(workbook), exam, subject, students)


def _norm_name(value):
    return re.sub(r'\s+', ' ', (value or '').strip().lower())


def _lookup_student(row, columns, students_by_id, students_by_roll, students_by_roll_name):
    """Find the exam student for this Excel row.

    After a re-import the Student ID often changes while Roll + Name stay the
    same. Prefer the current exam ID; if that ID is not on this roll, match
    Roll + Name so a teacher-filled old sheet still imports.
    """
    student_id = _cell_text(row[columns['id']]) if 'id' in columns and columns['id'] < len(row) else ''
    if student_id:
        student = students_by_id.get(student_id.lower())
        if student:
            return student, student_id
    roll_raw = _cell_text(row[columns['roll']]) if 'roll' in columns and columns['roll'] < len(row) else ''
    name = _cell_text(row[columns['name']]) if 'name' in columns and columns['name'] < len(row) else ''
    if roll_raw.isdigit() and name:
        matches = students_by_roll_name.get((int(roll_raw), _norm_name(name)), [])
        if len(matches) == 1:
            return matches[0], student_id or roll_raw
    if roll_raw.isdigit() and not student_id:
        matches = students_by_roll.get(int(roll_raw), [])
        if len(matches) == 1:
            return matches[0], roll_raw
        if len(matches) > 1:
            raise ValueError('roll number is shared by more than one student — use the ID column')
    return None, student_id or roll_raw


def _column_map(headers):
    mapped = {}
    for index, raw in enumerate(headers):
        key = _header_key(raw)
        if key in {'roll', 'roll id', 'roll no', 'roll number', 'roll_no'}:
            mapped.setdefault('roll', index)
        elif key in {'id', 'student id', 'student_id'}:
            mapped.setdefault('id', index)
        elif key in {'name', 'student name'}:
            mapped.setdefault('name', index)
        elif key in {'marks', 'total', 'marks obtained'}:
            mapped.setdefault('marks', index)
        elif key in {'subject code', 'subject', 'code'}:
            mapped.setdefault('subject_code', index)
        else:
            for part_key, aliases in PART_HEADER_ALIASES.items():
                if key in aliases:
                    mapped.setdefault(part_key, index)
                    break
    return mapped


def _parse_number(raw, label, maximum):
    if raw in (None, ''):
        return None
    text = _cell_text(raw)
    if text == '':
        return None
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f'{label} is not a number') from exc
    if number != number or number in (Decimal('Infinity'), Decimal('-Infinity')):
        raise ValueError(f'{label} is not a number')
    if number < 0 or number > maximum:
        raise ValueError(f'{label} must be between 0 and {maximum}')
    return number


def parse_subject_marks_sheet(sheet, exam, subject, students):
    """Return (validated_rows, skipped_count, errors).

    Each validated row is (student, defaults_dict) ready for ExamMark.update_or_create.
    """
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [_header_key(value) for value in header_row]
    columns = _column_map(header_row)
    if 'id' not in columns and 'roll' not in columns:
        raise ValueError('The first row must contain Roll and ID columns (or Student ID).')

    marks_config = get_subject_marks(exam, subject)
    parts = marks_config.parts
    students_by_id = {student.student_id.lower(): student for student in students if student.student_id}
    students_by_roll = defaultdict(list)
    students_by_roll_name = defaultdict(list)
    for student in students:
        if student.roll_no is not None:
            students_by_roll[int(student.roll_no)].append(student)
            students_by_roll_name[(int(student.roll_no), _norm_name(student.name))].append(student)

    is_legacy = headers[:3] == ['student id', 'subject code', 'marks']
    if not is_legacy:
        if parts:
            if not any(part['key'] in columns for part in parts):
                expected = ', '.join(EXCEL_PART_COLUMN[part['key']] for part in parts)
                raise ValueError(
                    f'The first row must contain {expected} columns '
                    f'(this subject is marked in parts, not a single total).'
                )
        elif 'marks' not in columns:
            raise ValueError('The first row must contain a Marks column.')
    validated, errors = [], []
    skipped = 0
    seen = set()

    for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(cell in (None, '') for cell in row):
            continue
        row = list(row)
        try:
            if is_legacy:
                student_id = _cell_text(row[0] if len(row) > 0 else '')
                subject_code = _cell_text(row[1] if len(row) > 1 else '')
                marks_raw = row[2] if len(row) > 2 else None
                if marks_raw in (None, ''):
                    skipped += 1
                    continue
                if subject_code and subject_code.lower() != subject.code.lower():
                    raise ValueError(
                        f'this file is for {subject.name} ({subject.code}), not {subject_code}'
                    )
                student = students_by_id.get(student_id.lower()) if student_id else None
                if not student:
                    skipped += 1
                    continue
                if student.pk in seen:
                    raise ValueError('duplicate student row')
                seen.add(student.pk)
                number = _parse_number(marks_raw, 'marks', marks_config.full_marks)
                defaults = {
                    'marks_obtained': number,
                    **{ExamMark.obtained_field(part['key']): None for part in MARK_PARTS},
                }
                validated.append((student, defaults))
                continue

            student, lookup = _lookup_student(
                row, columns, students_by_id, students_by_roll, students_by_roll_name,
            )
            if not student:
                skipped += 1
                continue
            if student.pk in seen:
                raise ValueError('duplicate student row')
            seen.add(student.pk)

            if parts:
                values = {}
                for part in parts:
                    index = columns.get(part['key'])
                    raw = row[index] if index is not None and index < len(row) else None
                    values[part['key']] = _parse_number(raw, part['label'], part['max_marks'])
                if all(value is None for value in values.values()):
                    skipped += 1
                    continue
                defaults = {
                    'marks_obtained': sum((v for v in values.values() if v is not None), Decimal('0')),
                }
                for part in MARK_PARTS:
                    defaults[ExamMark.obtained_field(part['key'])] = values.get(part['key'])
                validated.append((student, defaults))
            else:
                index = columns.get('marks')
                raw = row[index] if index is not None and index < len(row) else None
                if raw in (None, ''):
                    skipped += 1
                    continue
                number = _parse_number(raw, 'marks', marks_config.full_marks)
                defaults = {
                    'marks_obtained': number,
                    **{ExamMark.obtained_field(part['key']): None for part in MARK_PARTS},
                }
                validated.append((student, defaults))
        except (TypeError, ValueError) as exc:
            errors.append(f'Row {row_num}: {exc}')

    return validated, skipped, errors
