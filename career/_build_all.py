# -*- coding: utf-8 -*-
"""Build career package: .docx (Bangla font), index.html, all-in-one HTML, zip.
Run from career/:  python _build_all.py  (needs python-docx, polib not needed)"""
import os, re, html, zipfile, glob

HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)

DOCS = [p for p in sorted(glob.glob('*.md')) if not p.startswith('_')]
MARKET = sorted(glob.glob('market/*.md'))

def set_bangla_font(run):
    run.font.name = 'Nirmala UI'
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
    if rFonts is None:
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}' + attr.split(':')[1], 'Nirmala UI')

def add_rich(par, text):
    for j, part in enumerate(re.split(r'\*\*', text)):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = (j % 2 == 1)
        set_bangla_font(r)

def md_to_docx(md_path, docx_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Nirmala UI'; st.font.size = Pt(11)
    lines = open(md_path, encoding='utf-8').read().splitlines()
    i, n, in_code, table_rows = 0, len(lines), False, []
    NAVY = RGBColor(0x0F, 0x2D, 0x52)
    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        rows = [r for r in table_rows if not re.match(r'^\s*\|?\s*:?-{2,}', r.replace('|', ' ').replace(':', ' '))]
        parsed = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows if r.strip()]
        if parsed:
            t = doc.add_table(rows=len(parsed), cols=len(parsed[0]))
            t.style = 'Light Grid Accent 1'
            for ri, row in enumerate(parsed):
                for ci, cell in enumerate(row):
                    if ci < len(t.rows[ri].cells):
                        c = t.rows[ri].cells[ci]
                        c.text = ''
                        p = c.paragraphs[0]
                        add_rich(p, cell)
                        if ri == 0:
                            for r in p.runs:
                                r.bold = True
        table_rows = []
    while i < n:
        line = lines[i]
        if line.strip().startswith('```'):
            flush_table(); in_code = not in_code; i += 1; continue
        if in_code:
            p = doc.add_paragraph(); r = p.add_run(line); set_bangla_font(r); r.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(0)
            i += 1; continue
        if re.match(r'^\s*\|', line):
            table_rows.append(line); i += 1; continue
        flush_table()
        s = line.strip()
        if not s:
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            level = min(len(m.group(1)), 3)
            p = doc.add_heading('', level=level)
            add_rich(p, m.group(2).replace('**', ''))
            for r in p.runs:
                r.font.color.rgb = NAVY
            i += 1; continue
        if re.match(r'^(-{3,}|\*{3,})$', s):
            p = doc.add_paragraph('────────────────────')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1; continue
        if s.startswith('> '):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(24)
            add_rich(p, s[2:]); i += 1; continue
        m = re.match(r'^(\s*)[-*]\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Bullet'); add_rich(p, m.group(2)); i += 1; continue
        m = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph(style='List Number'); add_rich(p, m.group(2)); i += 1; continue
        p = doc.add_paragraph(); add_rich(p, s); i += 1
    flush_table()
    doc.save(docx_path)
    return docx_path

# ---------------- HTML ----------------
CSS = """
body{font-family:'Noto Sans Bengali','Hind Siliguri',sans-serif;margin:0;background:#F4F6FA;color:#1c2733;line-height:1.75}
.wrap{max-width:900px;margin:0 auto;padding:28px 20px 60px}
h1{color:#0F2D52;font-size:1.7rem;margin:34px 0 6px}h1:first-of-type{margin-top:0}
h2{color:#0F2D52;font-size:1.3rem;margin:28px 0 8px;border-bottom:2px solid #C9A227;padding-bottom:4px}
h3{color:#17406e;margin:20px 0 6px}
table{border-collapse:collapse;width:100%;margin:12px 0;background:#fff;font-size:.95rem}
th,td{border:1px solid #d7deea;padding:8px 10px;text-align:left;vertical-align:top}
th{background:#0F2D52;color:#fff}
tr:nth-child(even) td{background:#f7f9fd}
blockquote{margin:14px 0;padding:12px 18px;background:#fff8e6;border-left:4px solid #C9A227;border-radius:0 10px 10px 0}
code,pre{font-family:Consolas,monospace}code{background:#eef1f7;padding:1px 5px;border-radius:4px;font-size:.9em}
pre{background:#0F2D52;color:#e8edf5;padding:16px;border-radius:10px;overflow-x:auto;white-space:pre-wrap}
pre code{background:none;color:inherit}
a{color:#17406e}
.doc-toc{background:#fff;border:1px solid #e2e8f3;border-radius:12px;padding:14px 20px;margin:14px 0 26px}
.doc-toc a{display:block;padding:3px 0;text-decoration:none}
.topbar{position:sticky;top:0;background:#0F2D52;color:#fff;padding:10px 20px;display:flex;gap:16px;align-items:center;z-index:50}
.topbar a{color:#fff;text-decoration:none;font-size:.92rem}
.topbar .d{color:#C9A227;font-weight:600}
hr{border:none;border-top:1px solid #d7deea;margin:22px 0}
"""

def esc(t):
    return html.escape(t, quote=False)

def inline_md(t):
    t = esc(t)
    t = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', t)
    t = re.sub(r'`([^`]+)`', r'<code>\1</code>', t)
    t = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<a href="\2">\1</a>', t)
    return t

def md_to_html_fragment(md_text):
    lines = md_text.splitlines()
    out, i, n, in_code, table = [], 0, len(lines), False, []
    def flush_table():
        if table:
            rows = [r for r in table if not re.match(r'^\s*\|?[\s:|-]*-[-\s:|]*\|?\s*$', r)]
            parsed = [[c.strip() for c in r.strip().strip('|').split('|')] for r in rows if r.strip()]
            if parsed:
                out.append('<table>')
                for ri, row in enumerate(parsed):
                    tag = 'th' if ri == 0 else 'td'
                    out.append('<tr>' + ''.join(f'<{tag}>{inline_md(c)}</{tag}>' for c in row) + '</tr>')
                out.append('</table>')
            table.clear()
    while i < n:
        line = lines[i]
        if line.strip().startswith('```'):
            flush_table(); out.append('</pre>' if in_code else '<pre>'); in_code = not in_code; i += 1; continue
        if in_code:
            out.append(esc(line)); i += 1; continue
        if re.match(r'^\s*\|', line):
            table.append(line); i += 1; continue
        flush_table()
        s = line.strip()
        if not s:
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lv = min(len(m.group(1)), 4)
            out.append(f'<h{lv}>{inline_md(m.group(2))}</h{lv}>'); i += 1; continue
        if re.match(r'^(-{3,}|\*{3,})$', s):
            out.append('<hr>'); i += 1; continue
        if s.startswith('> '):
            out.append(f'<blockquote>{inline_md(s[2:])}</blockquote>'); i += 1; continue
        m = re.match(r'^[-*]\s+(.*)$', s)
        if m:
            if out and out[-1] == '</ul>':
                out.pop()
            elif not (out and out[-1].startswith('<li')):
                pass
            if not (out and (out[-1].startswith('<li') or out[-1] == '</li>')):
                pass
            if not (len(out) > 0 and out[-1] == ''):
                pass
            if not (out and out[-1] in ('<ul>', '</li>')):
                out.append('<ul>')
            out.append(f'<li>{inline_md(m.group(1))}</li>')
            i += 1
            if i < n and not re.match(r'^\s*[-*]\s+', lines[i]):
                out.append('</ul>')
            continue
        m = re.match(r'^\d+\.\s+(.*)$', s)
        if m:
            if not (out and out[-1] in ('<ol>', '</li>')):
                out.append('<ol>')
            out.append(f'<li>{inline_md(m.group(1))}</li>')
            i += 1
            if i < n and not re.match(r'^\s*\d+\.\s+', lines[i]):
                out.append('</ol>')
            continue
        out.append(f'<p>{inline_md(s)}</p>'); i += 1
    flush_table()
    return '\n'.join(out)

def doc_id(path):
    return os.path.splitext(os.path.basename(path))[0][:2]

DOC_META = {
    '00': ('শুরু এখান থেকে', '৩০ সেকেন্ডে পুরো পরিকল্পনা + আজকের ৩টি কাজ'),
    '01': ('বাস্তবতা', 'কেন ভালো প্রোডাক্ট একা জেতে না; অপারেটর-কাজকে মঞ্চ বানানো'),
    '02': ('নিরাপত্তার ১২ নিয়ম', 'যে নিয়ম মানলে আর কখনো হেনস্থা হবে না'),
    '03': ('৯০ দিনের পরিকল্পনা', 'দিন ১–৩০, ৩১–৬০, ৬১–৯০ — পরীক্ষা-বিভাগ সংস্করণ'),
    '04': ('প্রস্তাবপত্র (অধ্যক্ষ)', 'কপি-পেস্ট চিঠি: তথ্য-নিরাপত্তা/ব্যাকআপ'),
    '05': ('প্রমাণ-রেকর্ড', 'কার্যক্রম রেজিস্টার + মাসিক রিপোর্ট + ACR খসড়া'),
    '06': ('বাইরের বাজার — সারসংক্ষেপ', 'কেন বাইরের অফারই আসল লিভারেজ'),
    '07': ('কথোপকথন স্ক্রিপ্ট', 'অধ্যক্ষ, শিক্ষক, একাউন্টস, প্রভাবশালী ব্যক্তি — কী বলবেন'),
    '08': ('ট্রাস্ট ও চেয়ারম্যান কৌশল', '৫ প্রতিষ্ঠানের "এক ছাতা" পিচ + সম্প্রসারণের সিঁড়ি'),
    '09': ('চেয়ারম্যান-প্রপোজাল (৪ চিঠি)', 'কভারিং নোট + তথ্য-নিরাপত্তা + MIS পাইলট + দায়িত্বের বিবরণী'),
    '10': ('বাংলা UI রোডম্যাপ', 'কমিট ২a39683-এ যা হয়েছে + পরের ৪ সপ্তাহ'),
    '11': ('বাইরের বাজার প্যাকেজ', 'CV/লিংকডইন/ভিডিও কীভাবে ব্যবহার করবেন + অফার-কৌশল'),
}
MARKET_META = {
    'market/CV-English.md': ('CV (English)', 'বন্ধনী-ঘর ভরে নিলেই প্রস্তুত'),
    'market/linkedin-github.md': ('LinkedIn + GitHub সেটআপ', 'কপি-পেস্ট টেক্সট'),
    'market/demo-video-script.md': ('৩ মিনিটের ডেমো ভিডিও স্ক্রিপ্ট', 'শট-লিস্ট + বলার কথা'),
}

# --- individual docx ---
os.makedirs('word', exist_ok=True)
os.makedirs('market/word', exist_ok=True)
for md in DOCS:
    md_to_docx(md, f"word/{os.path.splitext(md)[0]}.docx")
for md in MARKET:
    md_to_docx(md, f"market/word/{os.path.splitext(os.path.basename(md))[0]}.docx")
md_to_docx('market/CV-English.md', 'word/CV-English.docx')
print('docx done:', len(DOCS) + len(MARKET) + 1)

# --- all-in-one html ---
parts = [f"<!DOCTYPE html><html lang='bn'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
         f"<title>ক্যারিয়ার প্যাকেজ — সব একসাথে</title><style>{CSS}</style></head><body>"]
parts.append("<div class='topbar'><span class='d'>ক্যারিয়ার প্যাকেজ</span><a href='index.html'>⟵ হাব</a></div><div class='wrap'>")
toc = ['<div class="doc-toc"><b>সূচি</b>']
for md in DOCS:
    k = doc_id(md); title = DOC_META[k][0]
    toc.append(f"<a href='#d{k}'>{k} — {title}</a>")
for md in MARKET:
    t, d = MARKET_META[md]
    toc.append(f"<a href='#d{esc(os.path.basename(md))}'>{t}</a>")
toc.append('</div>')
parts.extend(toc)
for md in DOCS:
    k = doc_id(md)
    txt = open(md, encoding='utf-8').read()
    parts.append(f"<section id='d{k}'>{md_to_html_fragment(txt)}</section><hr>")
for md in MARKET:
    txt = open(md, encoding='utf-8').read()
    parts.append(f"<section id='d{esc(os.path.basename(md))}'>{md_to_html_fragment(txt)}</section><hr>")
parts.append('</div></body></html>')
open('সব-একসাথে.html', 'w', encoding='utf-8').write('\n'.join(parts))

# --- index.html ---
cards = []
for md in DOCS:
    k = doc_id(md); t, d = DOC_META[k]
    cards.append(f"<a class='card' href='সব-একসাথে.html#d{k}'><span class='num'>{k}</span><span class='t'>{t}</span><span class='d'>{d}</span></a>")
mkt = ''.join(f"<a class='card' href='সব-একসাথে.html#d{esc(os.path.basename(md))}'><span class='num'>📦</span><span class='t'>{MARKET_META[md][0]}</span><span class='d'>{MARKET_META[md][1]}</span></a>" for md in MARKET)
INDEX = f"""<!DOCTYPE html><html lang='bn'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>
<title>ক্যারিয়ার হাব — স্কুল ম্যানেজমেন্ট সিস্টেম → প্রমোশন</title><style>
{CSS}
.hero{{background:linear-gradient(135deg,#0F2D52,#1b3d63);color:#fff;border-radius:16px;padding:26px 28px;margin-bottom:18px}}
.hero h1{{color:#fff;margin:0 0 8px;font-size:1.5rem}}.hero p{{margin:4px 0;opacity:.92}}
.green{{background:#e8f6ec;border:1px solid #bfe3c8;border-radius:12px;padding:16px 20px;margin:14px 0}}
.green a{{display:inline-block;background:#1c7c3c;color:#fff;padding:8px 16px;border-radius:8px;text-decoration:none;font-weight:600;margin:4px 6px 0 0}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:12px;margin:18px 0}}
.card{{background:#fff;border:1px solid #e2e8f3;border-radius:12px;padding:14px 16px;text-decoration:none;color:inherit;display:flex;flex-direction:column;gap:4px;transition:box-shadow .15s}}
.card:hover{{box-shadow:0 6px 18px rgba(15,45,82,.12)}}
.card .num{{color:#C9A227;font-weight:700;font-size:.85rem}}.card .t{{font-weight:700;color:#0F2D52}}.card .d{{font-size:.88rem;color:#5b6b82}}
h2.sec{{margin-top:26px}}
</style></head><body><div class='wrap'>
<div class='hero'><h1>সফটওয়্যার → স্বীকৃতি → প্রমোশন</h1>
<p><b>আপনি:</b> সহকারি আইটি অফিসার (পরীক্ষা বিভাগে কর্মরত) • <b>প্রতিষ্ঠান:</b> ৫টির ট্রাস্ট, সিদ্ধান্ত চেয়ারম্যান</p>
<p><b>হাতিয়ার:</b> নিজের তৈরি মাল্টি-ইনস্টিটিউশন সিস্টেম (৬২১ টেস্ট) + বাংলা UI (কমিট ২a39683) + ১২টি নথির ক্যারিয়ার-পরিকল্পনা</p></div>
<div class='green'><b>⬇ ডাউনলোড</b><br>
<a href='career-package.zip'>পুরো প্যাকেজ (.zip)</a>
<a href='word/CV-English.docx'>CV (Word)</a>
<a href='word/09-চেয়ারম্যান-প্রপোজাল.docx'>চেয়ারম্যান-চিঠি ৪টি (Word)</a></div>
<h2 class='sec'>নথি (এই পেজেই পড়ুন)</h2>
<div class='grid'>{''.join(cards)}</div>
<h2 class='sec'>বাইরের বাজারের টুল</h2>
<div class='grid'>{mkt}</div>
<p style='color:#5b6b82;font-size:.85rem'>সতর্কতা: অফিসের কম্পিউটারে ব্যক্তিগত ফাইল নয়। প্রিন্ট করে নিন বা নিজের ফোনে রাখুন।</p>
</div></body></html>"""
open('index.html', 'w', encoding='utf-8').write(INDEX)

# --- zip ---
with zipfile.ZipFile('career-package.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    z.write('index.html'); z.write('সব-একসাথে.html')
    for md in DOCS:
        z.write(md)
    for md in MARKET:
        z.write(md)
    for f in glob.glob('word/*.docx'):
        z.write(f)
    for f in glob.glob('market/word/*.docx'):
        if os.path.basename(f) != 'CV-English.docx':
            z.write(f, 'word/' + os.path.basename(f))
print('index + all-in-one + zip done')
print('zip size:', round(os.path.getsize('career-package.zip') / 1024), 'KB')
