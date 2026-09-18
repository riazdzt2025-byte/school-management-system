#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""career/*.md  ->  Word (.docx) with a Bengali-capable font.

Usage:  python3 career/_build_docs.py
Outputs: career/word/<name>.docx  and  career/word/সব-একসাথে.docx
"""
import os
import re
import glob

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BANGLA_FONT = "Nirmala UI"          # present on Windows 10+, renders Bengali well
FALLBACK = "Vrinda"

HERE = os.path.dirname(os.path.abspath(__file__))
WORD_DIR = os.path.join(HERE, "word")
os.makedirs(WORD_DIR, exist_ok=True)


def set_bangla_font(run, size=None, bold=None, color=None):
    run.font.name = BANGLA_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), BANGLA_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def style_base(doc):
    st = doc.styles['Normal']
    st.font.name = BANGLA_FONT
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attr), BANGLA_FONT)
    st.font.size = Pt(11)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')


def add_rich(par, text, size=11, bold=False, color=None):
    text = text.replace('\\', '')
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = par.add_run(part[2:-2])
            set_bangla_font(r, size, True, color)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = par.add_run(part[1:-1])
            set_bangla_font(r, size - 1, bold, (0x33, 0x33, 0x33))
        else:
            r = par.add_run(part)
            set_bangla_font(r, size, bold, color)


def is_table_row(line):
    s = line.strip()
    return s.startswith('|') and s.endswith('|') and s.count('|') >= 3


def split_row(line):
    cells = [c.strip() for c in line.strip().strip('|').split('|')]
    return cells


def is_separator(line):
    s = line.strip()
    if not s.startswith('|'):
        return False
    return bool(re.fullmatch(r'\|[\s:\-|]+\|', s))


def render_markdown(doc, md_text):
    lines = md_text.split('\n')
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(b if b.strip() else ' ')
                set_bangla_font(r, 10, False, (0x33, 0x33, 0x33))
            doc.add_paragraph()
            continue

        # table
        if is_table_row(stripped) and i + 1 < n and is_separator(lines[i + 1]):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < n and is_table_row(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            tbl = doc.add_table(rows=1, cols=len(header))
            tbl.style = 'Table Grid'
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = tbl.rows[0].cells
            for idx, htext in enumerate(header):
                hdr[idx].text = ''
                p = hdr[idx].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                add_rich(p, htext, size=10, bold=True)
            for row in rows:
                cells = tbl.add_row().cells
                for idx in range(len(header)):
                    val = row[idx] if idx < len(row) else ''
                    cells[idx].text = ''
                    p = cells[idx].paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    add_rich(p, val, size=10)
            doc.add_paragraph()
            continue

        # headings
        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            add_rich(p, stripped[4:], size=12, bold=True, color=(0x1F, 0x3B, 0x63))
            i += 1
            continue
        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            add_rich(p, stripped[3:], size=14, bold=True, color=(0x1F, 0x3B, 0x63))
            i += 1
            continue
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            add_rich(p, stripped[2:], size=17, bold=True, color=(0x14, 0x2A, 0x4A))
            i += 1
            continue

        if stripped.startswith('---') and len(set(stripped)) == 1:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run('─' * 60)
            set_bangla_font(r, 9, False, (0xAA, 0xAA, 0xAA))
            i += 1
            continue

        if stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            add_rich(p, stripped[2:], size=11, bold=False, color=(0x44, 0x44, 0x44))
            i += 1
            continue

        m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        if m:
            indent = len(m.group(1))
            marker = m.group(2)
            p = doc.add_paragraph(style='List Bullet' if marker in '-*' else 'List Number')
            if indent:
                p.paragraph_format.left_indent = Cm(0.6 + 0.6 * (indent // 2))
            p.paragraph_format.space_after = Pt(2)
            add_rich(p, m.group(3), size=11)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_rich(p, stripped, size=11)
        i += 1


def build(md_path, doc=None, keep_title=True):
    if doc is None:
        doc = Document()
        style_base(doc)
        for s in doc.sections:
            s.left_margin = s.right_margin = Cm(2.0)
            s.top_margin = s.bottom_margin = Cm(1.8)
    with open(md_path, encoding='utf-8') as fh:
        md_text = fh.read()
    if not keep_title:
        md_text = md_text.split('\n', 1)[1] if md_text.startswith('#') else md_text
    render_markdown(doc, md_text)
    return doc


def main():
    files = sorted(glob.glob(os.path.join(HERE, '*.md')))
    files = [f for f in files if not os.path.basename(f).startswith('_')]
    made = []

    # combined
    combined = Document()
    style_base(combined)
    for s in combined.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
    p = combined.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, 'ক্যারিয়ার পরিকল্পনা — সব একসাথে', size=20, bold=True, color=(0x14, 0x2A, 0x4A))
    p = combined.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(p, 'প্রমোশন আদায়ের ৯০ দিনের খসড়া · ১৮ সেপ্টেম্বর ২০২৬', size=11, color=(0x66, 0x66, 0x66))

    for f in files:
        build(f, combined)
        combined.add_page_break()
    combined_path = os.path.join(WORD_DIR, 'সব-একসাথে.docx')
    combined.save(combined_path)
    made.append(combined_path)

    for f in files:
        out = os.path.join(WORD_DIR, os.path.splitext(os.path.basename(f))[0] + '.docx')
        d = build(f)
        d.save(out)
        made.append(out)

    for m in made:
        print('OK', os.path.relpath(m, HERE), os.path.getsize(m), 'bytes')


if __name__ == '__main__':
    main()
